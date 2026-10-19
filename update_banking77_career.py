from pathlib import Path
import html,json,stat,zipfile
from docx import Document
ROOT=Path(__file__).resolve().parent
EXP=ROOT/'02-AI-Data-Annotation/experiments/banking77'

def run():
    m=json.loads((EXP/'results/metrics.json').read_text(encoding='utf-8'))['published_test_metrics']
    text=f"Public BANKING77: 10,003 training / 3,080 test queries, 77 intents; training-only Naive Bayes selection, {m['accuracy']:.2%} classifier accuracy and macro-F1 {m['macro_f1']:.4f}. Source labels credited to Casanueva et al. / PolyAI (CC BY 4.0). Synthetic 100-message QA baseline retained."
    title='BANKING77 Intent Classification & Annotation QA'
    target=ROOT/'career/resume.html';s=target.read_text(encoding='utf-8')
    import re
    s=re.sub(r'<h3>(?:Text Annotation &amp; Quality Control|Text Annotation & Quality Control|BANKING77 Intent Classification &amp; Annotation QA)</h3><p>.*?</p>',lambda _: '<h3>'+html.escape(title)+'</h3><p>'+html.escape(text)+'</p>',s,flags=re.S)
    s=s.replace('Personal portfolio projects built with AI assistance and synthetic data.','Personal portfolio projects built with AI assistance and synthetic/public-source data.')
    target.write_text(s,encoding='utf-8')
    path=ROOT/'career/resume.docx'
    with zipfile.ZipFile(path) as z:
        infos=z.infolist();assert len(infos)<10000;assert sum(i.file_size for i in infos)<256*1024*1024;assert z.testzip() is None
        for i in infos:
            assert not i.flag_bits&1 and not Path(i.filename).is_absolute() and '..' not in Path(i.filename).parts
            assert stat.S_IFMT(i.external_attr>>16)!=stat.S_IFLNK
        body=z.read('word/document.xml');assert b'<w:ins ' not in body and b'<w:del ' not in body
    doc=Document(path);changed=False
    for para in doc.paragraphs:
        if para.text.startswith('Text Annotation & Quality Control') or para.text.startswith(title):
            para.clear();para.add_run(title+' — ').bold=True;para.add_run(text);changed=True
        elif 'Personal portfolio projects built with AI assistance and synthetic data.' in para.text:
            for r in para.runs:r.text=r.text.replace('AI assistance and synthetic data.','AI assistance and synthetic/public-source data.')
    assert changed,'Expected project entry absent'
    doc.save(path);verified=Document(path);assert any(title in p.text for p in verified.paragraphs)
    print('Resume now distinguishes attributed real-data classifier results from synthetic QA.')

if __name__=='__main__':run()
