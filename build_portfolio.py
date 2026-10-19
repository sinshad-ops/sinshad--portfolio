from __future__ import annotations

import base64
import csv
import hashlib
import html
import importlib.metadata
import json
import math
import os
import random
import re
import shutil
import subprocess
import sys
import textwrap
import unicodedata
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


SEED = 20261008
TODAY = "2026-10-09"
AUTHOR = "Sinshad PS"
EMAIL = "sinshadsirajudheenpriv@gmail.com"
LINKEDIN = "https://www.linkedin.com/in/sinshad-ps-587677418/"
GITHUB = "https://github.com/sinshad-ops"

ROOT = Path(__file__).resolve().parent
random.seed(SEED)

PROJECTS = {
    "01-LLM-Evaluation": {
        "short": "llm",
        "title": "LLM Pairwise Evaluation",
        "source": "llm.json",
        "description": "Pairwise response evaluation against provisional references and a separate AI scoring pass.",
    },
    "02-AI-Data-Annotation": {
        "short": "annotation",
        "title": "Multi-task AI Data Annotation",
        "source": "annotation.json",
        "description": "Topic, sentiment, urgency, entity-span annotation and transparent keyword baselines.",
    },
    "03-Multilingual-AI-Evaluation": {
        "short": "multilingual",
        "title": "Malayalam Translation Evaluation",
        "source": "multilingual.json",
        "description": "Benchmark composition and automatic Unicode/integrity QC for Malayalam candidates.",
    },
    "04-Multimodal-AI-Evaluation": {
        "short": "multimodal",
        "title": "Synthetic Visual Evaluation",
        "source": None,
        "description": "Rendered geometric scenes, geometry-derived ground truth, and pixel-based visual baselines.",
    },
    "05-AI-Quality-Evaluation": {
        "short": "safety",
        "title": "AI Safety and Quality Triage",
        "source": "safety.json",
        "description": "Transparent lexical triage compared with provisional safety decisions.",
    },
    "06-Search-Relevance": {
        "short": "search",
        "title": "Search Relevance Evaluation",
        "source": "search.json",
        "description": "Lexical ranking evaluated with NDCG@4 and MRR on a fictional judged corpus.",
    },
}


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")


def write_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def jsonl_write(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def jsonl_read(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def csv_write(path: Path, rows: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields = sorted({k for r in rows for k in r}) if rows else ["id"]
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({
                k: json.dumps(v, ensure_ascii=False, sort_keys=True)
                if isinstance(v, (list, dict)) else v
                for k, v in row.items()
            })


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(65536), b""):
            h.update(block)
    return h.hexdigest()


def find_input(name: str) -> Path:
    candidates = [
        ROOT / name,
        ROOT / "source-data" / name,
        ROOT / "inputs" / name,
        ROOT.parent / name,
        ROOT.parent / "source-data" / name,
    ]
    for p in candidates:
        if p.is_file():
            return p
    matches = [
        p for p in ROOT.rglob(name)
        if "data" not in p.parts[-3:] and not any(part.startswith("0") and "-" in part for part in p.parts)
    ]
    if matches:
        return sorted(matches)[0]
    raise FileNotFoundError(f"Required source file not found: {name}")


def validate_unique(rows: list[dict[str, Any]], count: int, label: str) -> None:
    if len(rows) != count:
        raise ValueError(f"{label}: expected {count} rows, found {len(rows)}")
    ids = [str(r.get("id", "")) for r in rows]
    if any(not x for x in ids) or len(ids) != len(set(ids)):
        raise ValueError(f"{label}: IDs must be non-empty and unique")


def simple_font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for name in candidates:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            pass
    return ImageFont.load_default()


def normalize_inputs() -> dict[str, list[dict[str, Any]]]:
    llm = read_json(find_input("llm.json"))
    judge = read_json(find_input("llm-judge.json"))
    ann = read_json(find_input("annotation.json"))
    multi = read_json(find_input("multilingual.json"))
    safety = read_json(find_input("safety.json"))
    search = read_json(find_input("search.json"))

    validate_unique(llm, 50, "llm.json")
    validate_unique(judge, 50, "llm-judge.json")
    validate_unique(ann, 100, "annotation.json")
    validate_unique(multi, 50, "multilingual.json")
    validate_unique(safety, 100, "safety.json")
    validate_unique(search, 30, "search.json")

    # Normalize case only; preserve raw AI author/judge files unchanged.
    for rows, field in ((llm, "expected_preference"), (judge, "preference"), (multi, "expected_preference")):
        for row in rows:
            value = str(row[field]).strip().lower()
            row[field] = value.upper() if value in {"a", "b"} else value

    llm_ids = {str(x["id"]) for x in llm}
    judge_ids = {str(x["id"]) for x in judge}
    if llm_ids != judge_ids:
        raise ValueError("llm.json and llm-judge.json IDs do not match")

    for row in judge:
        if row.get("preference") not in {"A", "B", "tie"}:
            raise ValueError(f"Invalid judge preference for {row['id']}")
        for side in ("scores_a", "scores_b"):
            scores = row.get(side, {})
            for dim in ("accuracy", "instruction_following", "helpfulness", "clarity"):
                if not isinstance(scores.get(dim), int) or not 1 <= scores[dim] <= 5:
                    raise ValueError(f"Invalid {side}.{dim} for {row['id']}")

    for row in ann:
        if row.get("topic") not in {"billing", "account", "delivery", "technical", "refund", "general"}:
            raise ValueError(f"Invalid annotation topic: {row['id']}")
        if row.get("sentiment") not in {"positive", "neutral", "negative", "mixed"}:
            raise ValueError(f"Invalid sentiment: {row['id']}")
        if row.get("urgency") not in {"low", "medium", "high"}:
            raise ValueError(f"Invalid urgency: {row['id']}")

    for row in multi:
        if row.get("expected_preference") not in {"A", "B", "tie"}:
            raise ValueError(f"Invalid multilingual preference: {row['id']}")

    for row in safety:
        if row.get("expected_decision") not in {"accept", "revise", "escalate"}:
            raise ValueError(f"Invalid safety decision: {row['id']}")

    for row in search:
        results = row.get("results", [])
        if len(results) != 4:
            raise ValueError(f"Search query {row['id']} must have four results")
        if any(r.get("relevance") not in {0, 1, 2, 3} for r in results):
            raise ValueError(f"Invalid relevance for query {row['id']}")

    return {
        "llm": llm,
        "judge": judge,
        "annotation": ann,
        "multilingual": multi,
        "safety": safety,
        "search": search,
    }


def annotation_records(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    output, quarantine = [], []
    for index, row in enumerate(rows):
        item = dict(row)
        spans = []
        problems = []
        text = str(item.get("text", ""))
        for entity_index, ent in enumerate(item.get("entities", [])):
            value = str(ent.get("text", ""))
            label = ent.get("label")
            starts = [m.start() for m in re.finditer(re.escape(value), text)] if value else []
            if label not in {"ORDER_ID", "PRODUCT", "DATE", "ORG"}:
                problems.append(f"unknown_entity_label:{label}")
            if not starts:
                problems.append(f"entity_not_found:{value}")
                quarantine.append({
                    "id": item["id"],
                    "entity_index": entity_index,
                    "entity_text": value,
                    "reason": "reference entity is not an exact substring",
                })
                continue
            if len(starts) > 1:
                problems.append(f"repeated_entity_text:{value}:using_first")
            start = starts[0]
            spans.append({"text": value, "label": label, "start": start, "end": start + len(value)})
        item["entity_spans"] = spans
        item["split"] = "calibration" if index < 70 else "holdout"
        item["qc_warnings"] = problems
        output.append(item)
    return output, quarantine


def draw_visual_dataset(project_dir: Path) -> list[dict[str, Any]]:
    images_dir = project_dir / "data" / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    palette = {
        "teal": (29, 184, 166),
        "amber": (245, 166, 35),
        "blue": (67, 120, 220),
        "red": (222, 72, 82),
        "purple": (143, 92, 190),
        "green": (68, 165, 93),
    }
    bg = (242, 246, 250)
    records: list[dict[str, Any]] = []

    for i in range(120):
        rng = random.Random(SEED + i * 997)
        scene_id = f"scene-{i + 1:03d}"
        mode = i % 6
        im = Image.new("RGB", (480, 300), bg)
        draw = ImageDraw.Draw(im)
        draw.rounded_rectangle((8, 8, 471, 291), 14, fill=(255, 255, 255), outline=(195, 207, 219), width=2)
        boxes: list[dict[str, Any]] = []
        colors_used: list[str] = []
        question_type = ""
        question = ""
        answer: Any = None
        evidence = ""

        if mode in {0, 1, 2}:
            count = 3 + (i % 5)
            shape = ["circle", "square", "triangle"][i % 3]
            for j in range(count):
                cname = list(palette)[(i + j * 2) % len(palette)]
                colors_used.append(cname)
                x = 45 + (j % 4) * 102 + rng.randint(-5, 5)
                y = 65 + (j // 4) * 120 + rng.randint(-5, 5)
                size = 44 + ((i + j) % 3) * 8
                box = (x, y, x + size, y + size)
                if shape == "circle":
                    draw.ellipse(box, fill=palette[cname])
                elif shape == "square":
                    draw.rectangle(box, fill=palette[cname])
                else:
                    draw.polygon([(x + size // 2, y), (x, y + size), (x + size, y + size)], fill=palette[cname])
                boxes.append({"shape": shape, "color": cname, "box": list(box), "rotation": 0, "occluded": False})
            if mode == 0:
                question_type = "count"
                question = f"How many {shape}s are visible?"
                answer = count
                evidence = f"Rendering metadata contains {count} non-overlapping {shape} components."
            elif mode == 1:
                question_type = "dominant_color"
                dominant = Counter(colors_used).most_common(1)[0][0]
                question = "Which named color appears most often? If tied, use the leftmost object's color."
                max_count = max(Counter(colors_used).values())
                tied = {c for c, n in Counter(colors_used).items() if n == max_count}
                answer = next(c for c in colors_used if c in tied)
                evidence = f"Object color counts are {dict(Counter(colors_used))}; deterministic leftmost tie rule applies."
            else:
                question_type = "left_right"
                left = boxes[0]
                right = boxes[-1]
                question = f"Is the first {left['color']} shape left of the final {right['color']} shape?"
                answer = "yes" if left["box"][0] < right["box"][0] else "no"
                evidence = f"x coordinates are {left['box'][0]} and {right['box'][0]}."

        elif mode == 3:
            heights = [50 + ((i * 17 + j * 31) % 145) for j in range(4)]
            names = ["A", "B", "C", "D"]
            for j, height in enumerate(heights):
                x = 70 + j * 90
                cname = list(palette)[(i + j) % len(palette)]
                draw.rectangle((x, 245 - height, x + 48, 245), fill=palette[cname])
                draw.text((x + 17, 252), names[j], fill=(29, 42, 57), font=simple_font(18, True))
                boxes.append({"shape": "bar", "color": cname, "box": [x, 245 - height, x + 48, 245], "height": height})
            tallest = names[max(range(4), key=lambda j: heights[j])]
            question_type = "tallest_bar"
            question = "Which bar is tallest: A, B, C, or D?"
            answer = tallest
            evidence = f"Rendered bar heights are {dict(zip(names, heights))}."

        elif mode == 4:
            present = rng.choice(list(palette))
            requested = present if (i // 6) % 2 == 0 else rng.choice([c for c in palette if c != present])
            draw.rounded_rectangle((60, 50, 420, 250), 12, fill=(232, 239, 247), outline=(139, 155, 173), width=2)
            for j, shape in enumerate(["circle", "square", "triangle"]):
                x = 85 + j * 110 + rng.randint(-12, 12)
                y = rng.randint(90, 145)
                size = rng.randint(35, 65)
                box = [x, y, x + size, y + size]
                if shape == "circle":
                    draw.ellipse(box, fill=palette[present])
                elif shape == "square":
                    draw.rectangle(box, fill=palette[present])
                else:
                    draw.polygon([(x + size // 2, y), (x, y + size), (x + size, y + size)], fill=palette[present])
                boxes.append({"shape": shape, "color": present, "box": box})
            question_type = "absence"
            question = f"Is the color {requested} absent from the central panel?"
            answer = "no" if requested == present else "yes"
            evidence = f"All central objects use {present}; requested color is {requested}."

        else:
            warning_color = rng.choice(["amber", "red"])
            warning_left = (i // 6) % 2 == 0
            wx = rng.randint(95, 145) if warning_left else rng.randint(310, 355)
            bx = rng.randint(225, 270) if warning_left else rng.randint(70, 110)
            wy = rng.randint(85, 115)
            by = rng.randint(130, 175)
            bw = rng.randint(90, 125)
            draw.rounded_rectangle((40, 45, 440, 255), 14, fill=(247, 249, 252), outline=(171, 183, 196), width=2)
            draw.polygon([(wx, wy), (wx-32, wy+65), (wx+32, wy+65)], fill=palette[warning_color])
            draw.text((wx-8, wy+22), "!", fill=(255, 255, 255), font=simple_font(26, True))
            draw.rounded_rectangle((bx, by, bx+bw, by+38), 8, fill=palette["teal"])
            draw.rounded_rectangle((180, 215, 285, 237), 6, fill=(205, 215, 225))
            question_type = "multi_step"
            question = "Is the warning icon left of the teal primary button?"
            answer = "yes" if wx < bx+bw/2 else "no"
            evidence = f"Warning center x={wx}; primary-button center x={bx+bw/2}."
            boxes = [
                {"shape": "warning_triangle", "color": warning_color, "box": [wx-32, wy, wx+32, wy+65]},
                {"shape": "button", "color": "teal", "box": [bx, by, bx+bw, by+38]},
            ]

        path = images_dir / f"{scene_id}.png"
        im.save(path, optimize=False)
        wrong = {
            "yes": "no", "no": "yes", "A": "B", "B": "C", "C": "D", "D": "A"
        }.get(str(answer), None)
        if wrong is None:
            if isinstance(answer, int):
                wrong = answer + 1
            elif str(answer) in palette:
                wrong = next(c for c in palette if c != answer)
            else:
                wrong = "incorrect"

        base = {
            "scene_id": scene_id,
            "image": f"images/{scene_id}.png",
            "width": 480,
            "height": 300,
            "background_rgb": list(bg),
            "question_type": question_type,
            "question": question,
            "gold_answer": answer,
            "evidence": evidence,
            "objects": boxes,
            "palette": {k: list(v) for k, v in palette.items()},
            "ground_truth_source": "deterministic rendering geometry",
        }
        records.append({**base, "record_id": f"{scene_id}-correct", "answer_variant": "correct",
                        "candidate_answer": answer, "expected_error": "none",
                        "rationale": f"The candidate matches the geometry-derived answer: {evidence}"})
        records.append({**base, "record_id": f"{scene_id}-faulty", "answer_variant": "controlled_fault",
                        "candidate_answer": wrong, "expected_error": "controlled_wrong_answer",
                        "rationale": f"The candidate conflicts with the geometry-derived answer. {evidence}"})
    return records


ANALYSIS_SOURCE = r'''
from __future__ import annotations
import argparse, csv, json, math, re, sys, unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
PROJECTS = {
 "llm":"01-LLM-Evaluation","annotation":"02-AI-Data-Annotation",
 "multilingual":"03-Multilingual-AI-Evaluation","multimodal":"04-Multimodal-AI-Evaluation",
 "safety":"05-AI-Quality-Evaluation","search":"06-Search-Relevance"
}

def loadj(path):
 return [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]
def dump(path,obj):
 path.parent.mkdir(parents=True,exist_ok=True)
 with path.open("w",encoding="utf-8",newline="\n") as f:
  f.write(json.dumps(obj,ensure_ascii=False,indent=2,sort_keys=True)+"\n")
def csvout(path,rows):
 path.parent.mkdir(parents=True,exist_ok=True)
 fields=sorted({k for r in rows for k in r}) if rows else ["id"]
 with path.open("w",encoding="utf-8",newline="") as f:
  w=csv.DictWriter(f,fieldnames=fields);w.writeheader()
  for r in rows:w.writerow({k:json.dumps(v,ensure_ascii=False,sort_keys=True) if isinstance(v,(dict,list)) else v for k,v in r.items()})
def div(a,b): return a/b if b else 0.0
def confusion(gold,pred,labels):
 return {g:{p:sum(1 for a,b in zip(gold,pred) if a==g and b==p) for p in labels} for g in labels}
def class_metrics(gold,pred,labels):
 out={}
 for lab in labels:
  tp=sum(a==lab and b==lab for a,b in zip(gold,pred)); fp=sum(a!=lab and b==lab for a,b in zip(gold,pred)); fn=sum(a==lab and b!=lab for a,b in zip(gold,pred))
  pr=div(tp,tp+fp); rc=div(tp,tp+fn)
  out[lab]={"precision":pr,"recall":rc,"f1":div(2*pr*rc,pr+rc),"support":sum(a==lab for a in gold)}
 return out
def classification(gold,pred,labels):
 pc=class_metrics(gold,pred,labels)
 return {"accuracy":div(sum(a==b for a,b in zip(gold,pred)),len(gold)),"macro_f1":sum(x["f1"] for x in pc.values())/len(labels),"per_class":pc,"confusion":confusion(gold,pred,labels),"n":len(gold)}
def font(size,bold=False):
 for p in ["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf","C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"]:
  try:return ImageFont.truetype(p,size)
  except OSError:pass
 return ImageFont.load_default()
def chart(path,title,values):
 values=list(values.items())
 im=Image.new("RGB",(960,500),(9,22,38));d=ImageDraw.Draw(im)
 d.text((45,28),title,fill=(238,244,248),font=font(28,True))
 maxv=max([v for _,v in values]+[1])
 for i,(name,val) in enumerate(values):
  y=95+i*max(48,int(350/max(1,len(values))))
  d.text((45,y),str(name)[:28],fill=(207,220,230),font=font(17))
  x=290; w=int(590*val/maxv)
  d.rounded_rectangle((x,y,x+max(3,w),y+25),6,fill=(26,188,167))
  d.text((x+max(8,w)+10,y),f"{val:.3f}" if isinstance(val,float) else str(val),fill=(247,181,57),font=font(16,True))
 im.save(path)
def tokens(s): return re.findall(r"[a-z0-9]+",str(s).lower())
def flatten_conf(c): return [{"gold":g,"predicted":p,"count":n} for g,row in c.items() for p,n in row.items()]

def llm(p):
 pairs=loadj(p/"data"/"pairs.jsonl"); judges=loadj(p/"data"/"judge-decisions.jsonl"); jm={str(x["id"]):x for x in judges}
 dims=["accuracy","instruction_following","helpfulness","clarity"]; weights={"accuracy":.4,"instruction_following":.3,"helpfulness":.2,"clarity":.1}
 gold=[];pred=[];cats=defaultdict(lambda:[0,0]);dimvals=defaultdict(list);passes=[];queue=[]
 for x in pairs:
  j=jm[str(x["id"])]; g=x["expected_preference"]; q=j["preference"];gold.append(g);pred.append(q);cats[x["category"]][1]+=1;cats[x["category"]][0]+=g==q
  for side in ["a","b"]:
   s=j["scores_"+side]
   for d in dims:dimvals[d].append(s[d])
   weighted=sum(weights[d]*s[d] for d in dims)
   passed=weighted>=4 and s["accuracy"]>=4 and s["instruction_following"]>=4
   passes.append({"id":x["id"],"response":side.upper(),"weighted_score":weighted,"accuracy":s["accuracy"],"instruction_following":s["instruction_following"],"pass":passed})
  if g!=q:queue.append({"id":x["id"],"category":x["category"],"provisional_preference":g,"judge_preference":q,"reason":"provisional author reference and AI judge disagree","status":"manual_review_pending"})
 labels=["A","B","tie"]; cl=classification(gold,pred,labels)
 sens={}
 for t in [3.5,4.0,4.5]:
  sens[str(t)]=div(sum(r["weighted_score"]>=t and r["accuracy"]>=4 and r["instruction_following"]>=4 for r in passes),len(passes))
 metrics={"project":"llm","case_count":50,"response_count":100,"provenance":{"author_reference":"Strawberry fast model, family/version unspecified","separate_scoring_pass":"Strawberry fast model, family/version unspecified","human_gold":False,"timestamp":"2026-10-08"},"agreement_rate":cl["accuracy"],"confusion":cl["confusion"],"ties":{"provisional":gold.count("tie"),"judge":pred.count("tie")},"per_category":{k:{"agreement":div(v[0],v[1]),"n":v[1]} for k,v in cats.items()},"mean_scores":{d:sum(v)/len(v) for d,v in dimvals.items()},"weighted_formula":"0.4*accuracy + 0.3*instruction_following + 0.2*helpfulness + 0.1*clarity","pass_rule":"weighted>=4 and accuracy>=4 and instruction_following>=4","pass_rate":div(sum(x["pass"] for x in passes),len(passes)),"threshold_sensitivity":sens,"disagreements":len(queue)}
 csvout(p/"results"/"review_queue.csv",queue);csvout(p/"results"/"response_scores.csv",passes)
 chart(p/"results"/"chart.png","Provisional reference vs AI judge",{"agreement":metrics["agreement_rate"],"pass_rate":metrics["pass_rate"],"disagreement_rate":1-metrics["agreement_rate"]})
 return metrics

TOPIC={
 "billing":["bill","charged","charge","invoice","payment","card","fee"],
 "account":["account","login","password","sign in","profile","locked"],
 "delivery":["delivery","delivered","shipping","shipment","parcel","courier","arrive","tracking"],
 "technical":["error","bug","app","website","crash","screen","connection","software"],
 "refund":["refund","money back","reimburse","return","cancel"]
}
SENT={"positive":["thanks","thank you","great","happy","excellent","love"],"negative":["angry","bad","terrible","failed","wrong","unhappy","frustrated","not working","late"]}
HIGH=["urgent","immediately","asap","emergency","today","now","locked out"]; MED=["soon","please help","waiting","delay","problem","issue"]
def choose_topic(text):
 t=text.lower();scores={k:sum(t.count(w) for w in ws) for k,ws in TOPIC.items()};m=max(scores.values())
 return "general" if m==0 else next(k for k in TOPIC if scores[k]==m)
def choose_sent(text):
 t=text.lower();p=sum(w in t for w in SENT["positive"]);n=sum(w in t for w in SENT["negative"])
 return "mixed" if p and n else "positive" if p else "negative" if n else "neutral"
def choose_urg(text):
 t=text.lower();return "high" if any(w in t for w in HIGH) else "medium" if any(w in t for w in MED) else "low"
def baseline_entities(text):
 pats=[("ORDER_ID",r"\b(?:ORD[- ]?\d{3,}|#[A-Z0-9]{4,})\b"),("DATE",r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]* \d{1,2}(?:, \d{4})?)\b")]
 out=[]
 for lab,pat in pats:
  for m in re.finditer(pat,text,re.I):out.append({"text":m.group(),"label":lab,"start":m.start(),"end":m.end()})
 return out
def annotation(p):
 rows=loadj(p/"data"/"production_data.jsonl");preds=[];q=[]
 for r in rows:
  pred={"id":r["id"],"split":r["split"],"topic":choose_topic(r["text"]),"sentiment":choose_sent(r["text"]),"urgency":choose_urg(r["text"]),"entities":baseline_entities(r["text"])}
  preds.append(pred)
  if r.get("ambiguous") or r.get("qc_warnings"):q.append({"id":r["id"],"ambiguous":r.get("ambiguous"),"warnings":r.get("qc_warnings"),"status":"manual_review_pending"})
 csvout(p/"data"/"baseline_predictions.csv",preds)
 result={"project":"annotation","case_count":len(rows),"split":{"calibration":70,"holdout":30},"baseline":"frozen keyword dictionaries; labels are not consulted during prediction"}
 for split in ["all","holdout"]:
  ix=[i for i,r in enumerate(rows) if split=="all" or r["split"]=="holdout"]
  result[split]={}
  for field,labels in [("topic",["billing","account","delivery","technical","refund","general"]),("sentiment",["positive","neutral","negative","mixed"]),("urgency",["low","medium","high"])]:
   result[split][field]=classification([rows[i][field] for i in ix],[preds[i][field] for i in ix],labels)
  tp=fp=fn=0
  for i in ix:
   g={(e["start"],e["end"],e["label"]) for e in rows[i].get("entity_spans",[])}
   z={(e["start"],e["end"],e["label"]) for e in preds[i]["entities"]}
   tp+=len(g&z);fp+=len(z-g);fn+=len(g-z)
  result[split]["entity_span_exact"]={"precision":div(tp,tp+fp),"recall":div(tp,tp+fn),"f1":div(2*tp,2*tp+fp+fn),"tp":tp,"fp":fp,"fn":fn}
 corrupt=[];changed=0
 for i,r in enumerate(rows):
  lab=r["topic"]
  if i%8==0:
   labs=["billing","account","delivery","technical","refund","general"];lab=labs[(labs.index(lab)+1)%len(labs)];changed+=1
  corrupt.append({"id":r["id"],"reference_topic":r["topic"],"simulated_second_topic":lab,"is_controlled_error":lab!=r["topic"]})
 csvout(p/"results"/"simulated_second_labels.csv",corrupt)
 result["simulation_not_human"]={"controlled_errors":changed,"rate":changed/len(rows),"agreement":1-changed/len(rows),"warning":"SIMULATION NOT HUMAN; not annotator agreement"}
 result["qc"]={"manual_queue":len(q),"malformed_span_quarantine":sum(bool(r.get("qc_warnings")) for r in rows)}
 csvout(p/"results"/"review_queue.csv",q)
 chart(p/"results"/"chart.png","Holdout keyword baseline",{"topic_accuracy":result["holdout"]["topic"]["accuracy"],"sentiment_accuracy":result["holdout"]["sentiment"]["accuracy"],"urgency_accuracy":result["holdout"]["urgency"]["accuracy"],"entity_f1":result["holdout"]["entity_span_exact"]["f1"]})
 return result

def multilingual(p):
 rows=loadj(p/"data"/"pairs.jsonl");prefs=Counter();errors=Counter();domains=defaultdict(lambda:[0,0]);qc=[];total=0
 for r in rows:
  prefs[r["expected_preference"]]+=1
  for side in ["a","b"]:
   total+=1; txt=r[f"candidate_{side}_malayalam"]; flags=[]
   if unicodedata.normalize("NFC",txt)!=txt:flags.append("not_nfc")
   ratio=div(sum(0x0D00<=ord(c)<=0x0D7F for c in txt),sum(c.isalpha() for c in txt))
   if ratio<.5:flags.append("low_malayalam_script_ratio")
   engdigits=re.findall(r"\d+",r["english"])
   if engdigits and not all(x in txt for x in engdigits):flags.append("english_digit_not_preserved_review_spelled_or_local_digits")
   errs=r.get(f"error_{side}",[])
   if isinstance(errs,str):errs=[errs]
   errors.update(errs);domains[r["domain"]][1]+=1;domains[r["domain"]][0]+=any(e not in {"none", ""} for e in errs)
   qc.append({"id":r["id"],"candidate":side.upper(),"flags":";".join(flags),"script_ratio":ratio,"status":"manual_review_pending" if flags else "automatic_checks_passed_not_language_validated"})
 csvout(p/"results"/"review_queue.csv",[x for x in qc if x["flags"]])
 metrics={"project":"multilingual","pair_count":len(rows),"candidate_count":total,"benchmark_preference_counts":dict(prefs),"synthetic_error_distribution":dict(errors),"per_domain_composition_error_rate":{k:{"rate":div(v[0],v[1]),"candidate_n":v[1]} for k,v in domains.items()},"automatic_qc_flags":sum(bool(x["flags"]) for x in qc),"human_review":"Malayalam native human review required; no human scores or language accuracy reported","integrity_scope":"flag-only Unicode/script/digit checks, never definitive correctness"}
 chart(p/"results"/"chart.png","Benchmark composition: expected preference",dict(prefs))
 return metrics

def components(im,colors):
 pix=im.load();w,h=im.size;target={tuple(v):k for k,v in colors.items()};seen=set();counts=Counter()
 for y in range(h):
  for x in range(w):
   rgb=pix[x,y]
   if rgb not in target or (x,y) in seen:continue
   stack=[(x,y)];seen.add((x,y));n=0
   while stack:
    a,b=stack.pop();n+=1
    for c,d in ((a-1,b),(a+1,b),(a,b-1),(a,b+1)):
     if 0<=c<w and 0<=d<h and (c,d) not in seen and pix[c,d]==rgb:seen.add((c,d));stack.append((c,d))
   if n>=100:counts[target[rgb]]+=1
 return counts
def visual_predict(r,p):
 typ=r["question_type"]; im=Image.open(p/"data"/r["image"]).convert("RGB"); cols=components(im,r["palette"])
 if typ=="count":return sum(cols.values())
 if typ=="dominant_color":
  if not cols:return None
  m=max(cols.values()); tied={c for c,n in cols.items() if n==m}
  pix=im.load()
  for x in range(im.width):
   for y in range(im.height):
    for c in tied:
     if pix[x,y]==tuple(r["palette"][c]):return c
 if typ=="tallest_bar":
  pix=im.load(); xs=[94,184,274,364]; tops=[]
  bg=(255,255,255)
  allowed={tuple(v) for v in r["palette"].values()}
  for x in xs:
   ys=[y for y in range(im.height) if pix[x,y] in allowed]
   tops.append(min(ys) if ys else 999)
  return ["A","B","C","D"][tops.index(min(tops))]
 return None
def multimodal(p):
 allrows=loadj(p/"data"/"visual_answers.jsonl");rows=[r for r in allrows if r["answer_variant"]=="correct"];evals=[];correct=0;eligible=0
 for r in rows:
  pred=visual_predict(r,p)
  if pred is not None:eligible+=1;ok=str(pred)==str(r["gold_answer"]);correct+=ok
  else:ok=False
  evals.append({"scene_id":r["scene_id"],"question_type":r["question_type"],"gold":r["gold_answer"],"pixel_prediction":pred if pred is not None else "ABSTAIN","correct":ok if pred is not None else "","status":"automatic_pixel_evaluation" if pred is not None else "abstained_requires_review"})
 csvout(p/"results"/"pixel_predictions.csv",evals);csvout(p/"results"/"review_queue.csv",[x for x in evals if x["pixel_prediction"]=="ABSTAIN" or x["correct"] is False])
 by=defaultdict(lambda:[0,0])
 for x in evals:
  if x["pixel_prediction"]!="ABSTAIN":by[x["question_type"]][1]+=1;by[x["question_type"]][0]+=bool(x["correct"])
 metrics={"project":"multimodal","scene_count":len(rows),"answer_record_count":len(allrows),"synthetic_correct_records":len(rows),"controlled_fault_records":len(allrows)-len(rows),"pixel_baseline":{"eligible":eligible,"correct":correct,"accuracy":div(correct,eligible),"coverage":div(eligible,len(rows)),"abstentions":len(rows)-eligible,"per_type":{k:{"accuracy":div(v[0],v[1]),"n":v[1]} for k,v in by.items()}},"ground_truth":"constructed from rendering geometry; metadata is an oracle only for dataset creation, never read by pixel baseline","limitation":"Artificial geometric images are not evidence of performance on natural photos or a VLM."}
 chart(p/"results"/"chart.png","Pixel baseline",{"accuracy_on_eligible":metrics["pixel_baseline"]["accuracy"],"coverage":metrics["pixel_baseline"]["coverage"],"abstention_rate":1-metrics["pixel_baseline"]["coverage"]})
 return metrics

ESC=["weapon","self-harm","suicide","explosive","malware","credential","bypass","dox","threat","violence","illegal"]
REV=["medical","legal","financial","guarantee","always","never","personal data","password","unsafe","unverified","harassment","hate"]
def triage(prompt,response):
 t=(str(prompt)+" "+str(response)).lower()
 if any(w in t for w in ESC):return "escalate"
 if any(w in t for w in REV):return "revise"
 return "accept"
def safety(p):
 rows=loadj(p/"data"/"cases.jsonl");pred=[triage(r["prompt"],r["response"]) for r in rows];gold=[r["expected_decision"] for r in rows];labels=["accept","revise","escalate"];cl=classification(gold,pred,labels);q=[];cats=defaultdict(lambda:[0,0])
 for r,z in zip(rows,pred):
  cats[r["category"]][1]+=1;cats[r["category"]][0]+=z==r["expected_decision"]
  if z=="escalate" or z!=r["expected_decision"]:q.append({"id":r["id"],"category":r["category"],"expected":r["expected_decision"],"baseline":z,"severity":r["severity"],"reason":"escalation or provisional-label mismatch","status":"human_policy_review_pending"})
 fn=sum(g=="escalate" and z!="escalate" for g,z in zip(gold,pred));fp=sum(g!="escalate" and z=="escalate" for g,z in zip(gold,pred))
 csvout(p/"results"/"review_queue.csv",q)
 metrics={"project":"safety","case_count":len(rows),"baseline":"lexical heuristic; expected labels are not consulted","accuracy":cl["accuracy"],"macro_f1":cl["macro_f1"],"per_class":cl["per_class"],"confusion":cl["confusion"],"escalation_false_negatives":fn,"escalation_false_positives":fp,"category_breakdown":{k:{"accuracy":div(v[0],v[1]),"n":v[1]} for k,v in cats.items()},"risk_tradeoff":"Escalation false negatives receive priority; false positives expose over-refusal and reviewer cost.","scope":"Portfolio-owned provisional rubric, not employer policy and not a safety certification."}
 chart(p/"results"/"chart.png","Safety triage baseline",{"accuracy":cl["accuracy"],"macro_f1":cl["macro_f1"],"escalation_FN":fn,"escalation_FP":fp})
 return metrics

def dcg(rels):return sum((2**r-1)/math.log2(i+2) for i,r in enumerate(rels))
def search(p):
 qs=loadj(p/"data"/"queries.jsonl");nd=[];rr=[];intents=defaultdict(list);fails=[];rankrows=[]
 N=sum(len(q["results"]) for q in qs)
 df=Counter()
 for q in qs:
  for r in q["results"]:df.update(set(tokens(r["title"]+" "+r["snippet"])))
 for q in qs:
  qt=Counter(tokens(q["query"]));sc=[]
  for idx,r in enumerate(q["results"]):
   dt=Counter(tokens(r["title"]+" "+r["snippet"]))
   score=sum(qn*dt[t]*(math.log((N+1)/(df[t]+1))+1) for t,qn in qt.items())
   sc.append((score,idx,r))
  ranked=[x[2] for x in sorted(sc,key=lambda x:(-x[0],x[1]))];rels=[r["relevance"] for r in ranked];ideal=sorted([r["relevance"] for r in q["results"]],reverse=True)
  n=dcg(rels)/dcg(ideal) if dcg(ideal) else 0; m=next((1/(i+1) for i,r in enumerate(rels) if r>=2),0)
  nd.append(n);rr.append(m);intents[q["intent"]].append(n)
  for rank,r in enumerate(ranked,1):rankrows.append({"query_id":q["id"],"rank":rank,"result_id":r["id"],"relevance":r["relevance"],"query":q["query"]})
  if n<.8:fails.append({"id":q["id"],"query":q["query"],"intent":q["intent"],"locale":q["locale"],"ndcg_at_4":n,"status":"manual_relevance_review_pending"})
 csvout(p/"results"/"ranked_results.csv",rankrows);csvout(p/"results"/"review_queue.csv",fails)
 metrics={"project":"search","query_count":len(qs),"result_count":sum(len(q["results"]) for q in qs),"ndcg_at_4":sum(nd)/len(nd),"mrr_relevance_gte_2":sum(rr)/len(rr),"per_intent_ndcg":{k:sum(v)/len(v) for k,v in intents.items()},"failure_queue":len(fails),"ranking":"TF-IDF-style lexical query/title/snippet overlap; gold relevance is not used for ranking","scope":"Fictional corpus judgments, not live-search performance; locale and local intent require reviewer context."}
 chart(p/"results"/"chart.png","Search ranking quality",{"NDCG@4":metrics["ndcg_at_4"],"MRR":metrics["mrr_relevance_gte_2"],"failure_rate":len(fails)/len(qs)})
 return metrics

FUN={"llm":llm,"annotation":annotation,"multilingual":multilingual,"multimodal":multimodal,"safety":safety,"search":search}
def run(name):
 p=ROOT/PROJECTS[name];m=FUN[name](p);dump(p/"results"/"metrics.json",m);return m
if __name__=="__main__":
 a=argparse.ArgumentParser();a.add_argument("--project",choices=["all"]+list(PROJECTS),default="all");a.add_argument("--print",action="store_true");z=a.parse_args()
 out={k:run(k) for k in PROJECTS} if z.project=="all" else {z.project:run(z.project)}
 if z.print:print(json.dumps(out,ensure_ascii=False,indent=2,sort_keys=True))
'''


def prepare_projects(data: dict[str, list[dict[str, Any]]]) -> None:
    for dirname in PROJECTS:
        p = ROOT / dirname
        for sub in ("data", "src", "notebooks", "results", "qa"):
            (p / sub).mkdir(parents=True, exist_ok=True)

    p1 = ROOT / "01-LLM-Evaluation"
    jsonl_write(p1 / "data" / "pairs.jsonl", data["llm"])
    jsonl_write(p1 / "data" / "judge-decisions.jsonl", data["judge"])
    csv_write(p1 / "data" / "pairs.csv", data["llm"])
    csv_write(p1 / "data" / "judge-decisions.csv", data["judge"])

    p2 = ROOT / "02-AI-Data-Annotation"
    ann, quarantine = annotation_records(data["annotation"])
    jsonl_write(p2 / "data" / "production_data.jsonl", ann)
    csv_write(p2 / "data" / "production_data.csv", ann)
    csv_write(p2 / "qa" / "quarantined_spans.csv", quarantine,
              ["id", "entity_index", "entity_text", "reason"])

    p3 = ROOT / "03-Multilingual-AI-Evaluation"
    jsonl_write(p3 / "data" / "pairs.jsonl", data["multilingual"])
    csv_write(p3 / "data" / "pairs.csv", data["multilingual"])
    reviewer_rows = []
    for r in data["multilingual"]:
        for side in ("A", "B"):
            reviewer_rows.append({
                "id": r["id"], "candidate": side, "domain": r["domain"],
                "meaning_1_5": "", "grammar_1_5": "", "naturalness_1_5": "",
                "register_1_5": "", "reviewer_rationale": "", "reviewer_name": "",
                "review_date": "", "status": "blank_native_human_review_template",
            })
    csv_write(p3 / "qa" / "native_reviewer_template.csv", reviewer_rows)

    p4 = ROOT / "04-Multimodal-AI-Evaluation"
    visuals = draw_visual_dataset(p4)
    jsonl_write(p4 / "data" / "visual_answers.jsonl", visuals)
    csv_write(p4 / "data" / "visual_answers.csv", visuals)
    forms = [{
        "scene_id": r["scene_id"], "image": r["image"], "question": r["question"],
        "gold_answer": r["gold_answer"], "reviewer_answer": "", "evidence_region": "",
        "confidence_1_5": "", "notes": "", "status": "blank_visual_review_form",
    } for r in visuals if r["answer_variant"] == "correct"]
    jsonl_write(p4 / "qa" / "visual_evaluation_forms.jsonl", forms)
    create_contact_sheet(p4)

    p5 = ROOT / "05-AI-Quality-Evaluation"
    jsonl_write(p5 / "data" / "cases.jsonl", data["safety"])
    csv_write(p5 / "data" / "cases.csv", data["safety"])

    p6 = ROOT / "06-Search-Relevance"
    jsonl_write(p6 / "data" / "queries.jsonl", data["search"])
    csv_write(p6 / "data" / "queries.csv", data["search"])
    flat = []
    for q in data["search"]:
        for r in q["results"]:
            flat.append({
                "query_id": q["id"], "query": q["query"], "intent": q["intent"],
                "locale": q["locale"], **r
            })
    jsonl_write(p6 / "data" / "result_judgments.jsonl", flat)
    csv_write(p6 / "data" / "result_judgments.csv", flat)


def create_contact_sheet(p4: Path) -> None:
    sheet = Image.new("RGB", (1000, 720), (9, 22, 38))
    draw = ImageDraw.Draw(sheet)
    draw.text((30, 18), "Visual Evaluation Evidence Contact Sheet", fill=(239, 245, 249),
              font=simple_font(27, True))
    for i in range(12):
        source = p4 / "data" / "images" / f"scene-{i + 1:03d}.png"
        image = Image.open(source).convert("RGB")
        image.thumbnail((220, 140))
        x = 25 + (i % 4) * 245
        y = 70 + (i // 4) * 210
        sheet.paste(image, (x, y))
        draw.text((x, y + 146), f"scene-{i + 1:03d}", fill=(247, 181, 57), font=simple_font(15, True))
    sheet.save(p4 / "qa" / "rendered-evidence-screenshot.png")


def run_analysis() -> None:
    write_text(ROOT / "analysis.py", ANALYSIS_SOURCE)
    for dirname, meta in PROJECTS.items():
        thin = f'''"""Standalone entry point for {meta["title"]}."""
from pathlib import Path
import subprocess
import sys

if __name__ == "__main__":
    root = Path(__file__).resolve().parents[2]
    raise SystemExit(subprocess.call(
        [sys.executable, str(root / "analysis.py"), "--project", "{meta["short"]}", "--print"]
    ))
'''
        write_text(ROOT / dirname / "src" / "analyze.py", thin)
    subprocess.run([sys.executable, str(ROOT / "analysis.py"), "--project", "all"], cwd=ROOT, check=True)


def data_card(dirname: str, metrics: dict[str, Any]) -> str:
    title = PROJECTS[dirname]["title"]
    count = metrics.get("case_count", metrics.get("pair_count", metrics.get("scene_count",
            metrics.get("query_count", 0))))
    special = {
        "01-LLM-Evaluation": "50 pairwise prompts and 100 candidate responses; full judge inputs and raw decisions retained.",
        "02-AI-Data-Annotation": "100 unique customer-support-style messages; first 70 IDs are calibration and final 30 are holdout.",
        "03-Multilingual-AI-Evaluation": "50 English–Malayalam pairs and 100 candidates. References are provisional AI linguistic references.",
        "04-Multimodal-AI-Evaluation": "120 rendered PNG scenes and 240 synthetic correct/faulty answer records.",
        "05-AI-Quality-Evaluation": "100 abstract safety/quality cases without harmful operational detail.",
        "06-Search-Relevance": "30 queries with four judged fictional results each, for 120 result judgments.",
    }[dirname]
    return f"""# Data Card — {title}

## Summary
{special}

## Intended use
Local portfolio demonstration of reproducible evaluation, QA queues, metric computation, and evidence review. It is not production training data, employer work, or a validated benchmark.

## Provenance
The supplied source arrays and generated artifacts are AI-assisted synthetic personal-project data authored for {AUTHOR}. LLM, multilingual, safety, annotation, and search labels are provisional unless explicitly described as rendering-derived ground truth. No independent human validation has been completed.

## Size
Primary case count reported by the pipeline: **{count}**.

## Processing
Source arrays are normalized into UTF-8 JSONL and CSV. Nested values in CSV are JSON-encoded. Source IDs are preserved. Deterministic build seed: `{SEED}`.

## Risks and limitations
Synthetic composition can make patterns easier than real operations. Reviewer demographics, adjudication, inter-annotator agreement, natural distribution shift, and external validity are absent. Sensitive production deployment requires qualified human review.

## License
Code: MIT. Synthetic data: CC BY 4.0; see repository `DATA_LICENSE.md`.
"""


def project_docs(dirname: str, metrics: dict[str, Any]) -> None:
    p = ROOT / dirname
    meta = PROJECTS[dirname]
    title = meta["title"]
    short = meta["short"]

    specifics = {
        "llm": {
            "framework": "Compare the held-back provisional author preference with a separate AI scoring pass preference. Compute a three-class confusion matrix, agreement, category slices, score means, threshold sensitivity, and response-level pass rules.",
            "formula": "Weighted score = 0.4×accuracy + 0.3×instruction following + 0.2×helpfulness + 0.1×clarity. Pass requires weighted score ≥4, accuracy ≥4, and instruction following ≥4.",
            "limits": "Neither the provisional author preference nor AI judge is human gold. Both calls use the Strawberry fast tier; underlying model identity/version and temperature were not supplied. Independence of model family is not established. Agreement is not model accuracy or human kappa.",
        },
        "annotation": {
            "framework": "Frozen keyword dictionaries predict topic, sentiment, and urgency without reading reference labels. Regex extracts only high-confidence order/date entities. The first 70 IDs are calibration and final 30 are holdout.",
            "formula": "Accuracy and macro-F1 are computed separately per task. Entity precision/recall/F1 requires exact start, end, and label match.",
            "limits": "Keywords are brittle and entity coverage is intentionally narrow. The controlled second-label corruption is SIMULATION NOT HUMAN and cannot establish annotator agreement.",
        },
        "multilingual": {
            "framework": "Summarize benchmark preference/error composition, then run NFC, Malayalam-codepoint ratio, and digit-preservation flag checks.",
            "formula": "Composition error rate counts supplied synthetic error tags. Automatic QC flags are integrity indicators, not linguistic scores.",
            "limits": "Malayalam native human review is pending. Digit translation may legitimately use Malayalam digits or spelled numbers; flags are never definitive language correctness.",
        },
        "multimodal": {
            "framework": "Render deterministic scenes, derive ground truth from geometry, create controlled correct/faulty answers, and evaluate eligible tasks with RGB connected-component or bar-pixel scanning.",
            "formula": "Pixel accuracy = correct eligible predictions / eligible predictions. Coverage = eligible predictions / all scenes; unsupported question types abstain.",
            "limits": "These are artificial images, not natural photos. Metadata is an oracle for data construction only and is not read by the pixel baseline. This is not a VLM evaluation.",
        },
        "safety": {
            "framework": "Apply transparent escalation/revision lexical rules without consulting expected decisions; compare against provisional project-owned labels and route escalations or mismatches to human review.",
            "formula": "Report accuracy, macro-F1, confusion, category slices, and escalation false negatives/positives. Escalation false negatives receive review priority.",
            "limits": "The rubric is not an employer policy and the results are not a safety certification. Lexical matching can both miss risks and over-refuse benign content.",
        },
        "search": {
            "framework": "Rank each query’s four fictional results by TF-IDF-style lexical overlap over query, title, and snippet without using relevance labels.",
            "formula": "NDCG@4 uses gain 2^relevance−1 and log2 rank discount. MRR uses the first result with relevance ≥2.",
            "limits": "Judgments describe a fictional corpus, not live search. Lexical matching misses semantic, domain, locale, freshness, and local-intent nuance.",
        },
    }[short]

    headline_metric = {
        "llm": lambda: f"AI-reference agreement: {metrics['agreement_rate']:.1%}",
        "annotation": lambda: f"Holdout topic accuracy: {metrics['holdout']['topic']['accuracy']:.1%}",
        "multilingual": lambda: f"Automatic QC flags: {metrics['automatic_qc_flags']} of {metrics['candidate_count']}",
        "multimodal": lambda: f"Pixel baseline: {metrics['pixel_baseline']['accuracy']:.1%} on {metrics['pixel_baseline']['coverage']:.1%} coverage",
        "safety": lambda: f"Heuristic macro-F1: {metrics['macro_f1']:.3f}",
        "search": lambda: f"NDCG@4: {metrics['ndcg_at_4']:.3f}; MRR: {metrics['mrr_relevance_gte_2']:.3f}",
    }[short]()

    readme = f"""# {title}

> AI-assisted synthetic personal portfolio project by {AUTHOR}; not employment work. Human validation is pending unless explicitly stated otherwise.

## Overview
{meta["description"]}

## Problem
Evaluation pipelines need auditable inputs, explicit rubrics, repeatable metrics, evidence-linked review queues, and honest uncertainty boundaries.

## Objective
Build a fully local workflow that preserves raw evidence, runs a baseline independent of evaluation labels where applicable, computes real metrics, and exposes failures for review.

## Dataset
See [`data/DATA_CARD.md`](data/DATA_CARD.md). All records are UTF-8 and available as JSONL and CSV.

## Evaluation framework
{specifics["framework"]}

## Scoring
{specifics["formula"]}

## Tools
Python standard library and Pillow. Inputs and outputs remain local; no network service is required.

## Implementation
- `src/analyze.py` runs this project independently.
- Root `analysis.py --project {short}` contains reusable metric and plotting logic.
- `notebooks/analysis.ipynb` contains executable code and captured output from the real metrics file.
- `results/review_queue.csv` makes disagreements, failures, escalations, or abstentions visible.
- `qa/worked-cases.md` connects evidence, rubric, QC, and decisions.

## Results
**{headline_metric}**

The complete machine-readable output is [`results/metrics.json`](results/metrics.json), with interpretation in [`results/analysis.md`](results/analysis.md).

![Result chart](results/chart.png)

## Findings
The pipeline produces measurable baseline behavior and a concrete review queue. A high metric, where present, measures only agreement with this project’s provisional labels or constructed rendering ground truth.

## Limitations
{specifics["limits"]}

## Future work
Run blinded qualified human review, adjudicate disagreements, document reviewer eligibility, test robustness under distribution shift, and revise rules only on calibration data before re-evaluating holdout data.

## Reproduction
From repository root:

`python analysis.py --project {short}`

or:

`python {dirname}/src/analyze.py`

Run validation with:

`python -m unittest discover -s tests`

Deterministic seed: `{SEED}`. Build date: {TODAY}.
"""
    write_text(p / "README.md", readme)
    write_text(p / "data" / "DATA_CARD.md", data_card(dirname, metrics))

    rubric = f"""# Rubric — {title}

## Decision principles
1. Inspect the source evidence before labels or baseline output where blinding is possible.
2. Apply the written criteria consistently; do not infer missing facts.
3. Record a decision, confidence, evidence quote or region, and concise rationale.
4. Escalate ambiguity, policy uncertainty, malformed evidence, or material disagreement.
5. Never convert automatic QC flags into unreviewed claims of semantic correctness.

## Project scoring
{specifics["formula"]}

## Resolution
A reviewer may uphold the reference, uphold the baseline/candidate, revise both, or mark insufficient evidence. Material disagreements require a second review and adjudication. Human-review fields remain blank until a real reviewer completes them.
"""
    guidelines = f"""# Guidelines — {title}

## Workflow
- Confirm record ID and source completeness.
- Read the source and candidate output.
- Apply the rubric without consulting unrelated cases.
- Capture explicit evidence, not intuition alone.
- Run or inspect automated QC.
- Choose a decision and reason.
- Route unresolved cases to `results/review_queue.csv`.

## Frozen baseline rule
Baseline dictionaries and thresholds are documented in root `analysis.py`. Evaluation labels are not consulted during prediction for annotation, safety, search, or pixel baselines.

## Human review status
No completed human validation is represented. Empty reviewer templates are intentional. Do not backfill ratings from provisional AI labels.
"""
    methodology = f"""# Methodology — {title}

## Design
{specifics["framework"]}

## Computation
{specifics["formula"]}

## Reproducibility
The build uses seed `{SEED}`, stable ID order, local files, deterministic drawing, and JSON metrics. Re-running root `analysis.py` recreates metrics, queues, and charts.

## QA controls
Schema checks, unique-ID checks, label-range checks, image pixel checks, entity span validation, deterministic reproduction tests, and output-consistency tests are run automatically.

## Interpretation boundary
{specifics["limits"]}
"""
    write_text(p / "rubrics.md", rubric)
    write_text(p / "guidelines.md", guidelines)
    write_text(p / "methodology.md", methodology)
    write_text(p / "results" / "analysis.md", result_analysis(short, metrics, specifics))
    write_text(p / "qa" / "worked-cases.md", worked_cases(dirname))


def result_analysis(short: str, metrics: dict[str, Any], specifics: dict[str, str]) -> str:
    pretty = json.dumps(metrics, ensure_ascii=False, indent=2, sort_keys=True)
    return f"""# Results Analysis

## Scope
These numbers are generated from the repository data by `analysis.py`; they are not manually entered.

## Primary interpretation
{specifics["framework"]}

## Important boundary
{specifics["limits"]}

## Decision use
Metrics identify patterns and prioritize review. They do not replace human adjudication, prove production readiness, or establish the author’s personal accuracy.

## Machine-readable snapshot
```json
{pretty}
```
"""


def sample_rows(dirname: str) -> list[dict[str, Any]]:
    p = ROOT / dirname
    short = PROJECTS[dirname]["short"]
    if short == "llm":
        rows = jsonl_read(p / "data" / "pairs.jsonl")
        judges = {str(x["id"]): x for x in jsonl_read(p / "data" / "judge-decisions.jsonl")}
        return [{
            "id": r["id"], "input": r["prompt"] + "\n\nSource context:\n" + r["source_context"], "output": f"A: {r['candidate_a']}\nB: {r['candidate_b']}",
            "rubric": "Pairwise accuracy, instruction following, helpfulness, clarity",
            "qc": f"Provisional={r['expected_preference']}; AI judge={judges[str(r['id'])]['preference']}",
            "decision": judges[str(r["id"])]["preference"],
            "evidence": judges[str(r["id"])].get("rationale", ""),
        } for r in rows[:3]]
    if short == "annotation":
        rows = jsonl_read(p / "data" / "production_data.jsonl")
        return [{
            "id": r["id"], "input": r["text"],
            "output": f"topic={r['topic']}; sentiment={r['sentiment']}; urgency={r['urgency']}",
            "rubric": "Exact categorical labels plus exact entity spans",
            "qc": "; ".join(r.get("qc_warnings", [])) or "No structural warning",
            "decision": "accept reference for provisional benchmark",
            "evidence": r.get("rationale", ""),
        } for r in rows[:3]]
    if short == "multilingual":
        rows = jsonl_read(p / "data" / "pairs.jsonl")
        return [{
            "id": r["id"], "input": r["english"],
            "output": f"A: {r['candidate_a_malayalam']}\nB: {r['candidate_b_malayalam']}",
            "rubric": "Meaning, grammar, naturalness, register; native review pending",
            "qc": f"Expected preference={r['expected_preference']}; automatic checks are flag-only",
            "decision": "manual Malayalam review pending",
            "evidence": r.get("rationale_english", ""),
        } for r in rows[:3]]
    if short == "multimodal":
        rows = [x for x in jsonl_read(p / "data" / "visual_answers.jsonl") if x["answer_variant"] == "correct"]
        return [{
            "id": r["scene_id"], "input": r["question"], "output": str(r["candidate_answer"]),
            "image": "data:image/png;base64," + base64.b64encode((p / "data" / r["image"]).read_bytes()).decode("ascii"),
            "rubric": f"Question type: {r['question_type']}; exact answer",
            "qc": "Image exists and contains non-background pixels",
            "decision": f"gold={r['gold_answer']}",
            "evidence": r["evidence"],
        } for r in rows[:3]]
    if short == "safety":
        rows = jsonl_read(p / "data" / "cases.jsonl")
        return [{
            "id": r["id"], "input": r["prompt"] + "\n\nSource context:\n" + r["source_context"], "output": r["response"],
            "rubric": f"Owned policy triage; severity={r['severity']}",
            "qc": f"Expected error={r['error_type']}",
            "decision": r["expected_decision"],
            "evidence": r["rationale"],
        } for r in rows[:3]]
    rows = jsonl_read(p / "data" / "queries.jsonl")
    return [{
        "id": r["id"], "input": r["query"],
        "output": "\n".join(f"{x['title']} — {x['snippet']} — relevance {x['relevance']}" for x in r["results"]),
        "rubric": "0–3 relevance; NDCG gain 2^rel−1",
        "qc": f"intent={r['intent']}; locale={r['locale']}",
        "decision": "compare lexical order with ideal judged order",
        "evidence": r["reference_rationale"],
    } for r in rows[:3]]


def worked_cases(dirname: str) -> str:
    title = PROJECTS[dirname]["title"]
    sections = [f"# Worked Cases — {title}",
                "\nThese cases use actual repository records. Decisions are provisional or geometry-derived as stated; they are not completed human reviews.\n"]
    for i, row in enumerate(sample_rows(dirname), 1):
        sections.append(f"""## Case {i}: `{row['id']}`
**Input/evidence:** {row['input']}

**Candidate/output:** {row['output']}

**Rubric applied:** {row['rubric']}

**QC:** {row['qc']}

**Decision:** {row['decision']}

**Explicit evidence/rationale:** {row['evidence']}
""")
    return "\n".join(sections)


def create_notebooks() -> None:
    for dirname, meta in PROJECTS.items():
        metrics = read_json(ROOT / dirname / "results" / "metrics.json")
        output_text = json.dumps(metrics, ensure_ascii=False, indent=2, sort_keys=True)
        code = (
            "import json, subprocess, sys\n"
            "from pathlib import Path\n"
            "project = Path.cwd().parent if Path.cwd().name == 'notebooks' else Path.cwd()\n"
            "project = project.resolve()\n"
            f"subprocess.run([sys.executable, str(project.parent / 'analysis.py'), '--project', '{meta['short']}'], check=True, stdout=subprocess.DEVNULL)\n"
            "metrics = json.loads((project / 'results' / 'metrics.json').read_text(encoding='utf-8'))\n"
            "print(json.dumps(metrics, ensure_ascii=False, indent=2, sort_keys=True))"
        )
        execution = subprocess.run([sys.executable, "-c", code], cwd=ROOT / dirname, capture_output=True, text=True, encoding="utf-8", check=True)
        output_text = execution.stdout.rstrip("\n")
        nb = {
            "cells": [
                {
                    "cell_type": "markdown",
                    "metadata": {},
                    "source": [
                        f"# {meta['title']}\n",
                        "Executed local summary. Re-run root `analysis.py` before this notebook to recreate metrics and charts."
                    ],
                },
                {
                    "cell_type": "code",
                    "execution_count": 1,
                    "metadata": {},
                    "source": code.splitlines(True),
                    "outputs": [{
                        "name": "stdout",
                        "output_type": "stream",
                        "text": [output_text + "\n"],
                    }],
                },
            ],
            "metadata": {
                "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
                "language_info": {"name": "python", "version": f"{sys.version_info.major}.{sys.version_info.minor}"},
            },
            "nbformat": 4,
            "nbformat_minor": 5,
        }
        write_json(ROOT / dirname / "notebooks" / "analysis.ipynb", nb)


def requirements() -> None:
    packages = []
    for dist in ("Pillow", "python-docx"):
        try:
            packages.append(f"{dist}=={importlib.metadata.version(dist)}")
        except importlib.metadata.PackageNotFoundError:
            raise RuntimeError(f"Required installed package missing: {dist}")
    write_text(ROOT / "requirements.txt", "\n".join(packages))


def create_manifest() -> dict[str, Any]:
    files = []
    project_counts = {}
    for dirname in PROJECTS:
        p = ROOT / dirname
        metrics = read_json(p / "results" / "metrics.json")
        project_counts[dirname] = {
            k: v for k, v in metrics.items()
            if k.endswith("_count") or k in {"case_count", "response_count", "pair_count",
                                             "candidate_count", "scene_count", "answer_record_count",
                                             "query_count", "result_count"}
        }
        for file in sorted((p / "data").rglob("*")):
            if file.is_file():
                files.append({
                    "path": file.relative_to(ROOT).as_posix(),
                    "bytes": file.stat().st_size,
                    "sha256": sha256(file),
                })
    manifest = {
        "generated": TODAY,
        "seed": SEED,
        "algorithm": "SHA-256",
        "project_counts": project_counts,
        "files": files,
    }
    write_json(ROOT / "DATA_MANIFEST.json", manifest)
    return manifest


def root_docs(metrics: dict[str, dict[str, Any]]) -> None:
    m1, m2, m3, m4, m5, m6 = (
        metrics["01-LLM-Evaluation"], metrics["02-AI-Data-Annotation"],
        metrics["03-Multilingual-AI-Evaluation"], metrics["04-Multimodal-AI-Evaluation"],
        metrics["05-AI-Quality-Evaluation"], metrics["06-Search-Relevance"],
    )
    readme = f"""# AI Generalist Evaluation Portfolio — {AUTHOR}

> **AI-assisted synthetic personal projects, not employment work.** Built locally to demonstrate reproducible evaluation pipelines, evidence review, QA controls, and honest limitations. Human validation is pending except for deterministic rendering-derived visual ground truth.

## Recruiter 30-second path
1. Open the offline [`website/index.html`](website/index.html).
2. Inspect one project’s `results/metrics.json`, review queue, and worked cases.
3. Run `python analysis.py --project all`.
4. Run `python -m unittest discover -s tests`.
5. Read [`HUMAN_REVIEW.md`](HUMAN_REVIEW.md) and [`project-defensibility-guide.md`](project-defensibility-guide.md).

## Project cards

| Project | Actual scale | Measured result | What it means |
|---|---:|---:|---|
| [LLM Evaluation](01-LLM-Evaluation/) | 50 pairs / 100 responses | {m1['agreement_rate']:.1%} agreement | Agreement between provisional AI-authored preferences and a separate AI scoring pass; not human accuracy |
| [Data Annotation](02-AI-Data-Annotation/) | 100 messages | {m2['holdout']['topic']['accuracy']:.1%} holdout topic accuracy | Frozen keyword baseline against provisional references |
| [Malayalam Evaluation](03-Multilingual-AI-Evaluation/) | 50 pairs / 100 candidates | {m3['automatic_qc_flags']} QC flags | Benchmark composition and automatic integrity QC; no validated language accuracy |
| [Multimodal Evaluation](04-Multimodal-AI-Evaluation/) | 120 images / 240 answers | {m4['pixel_baseline']['accuracy']:.1%} eligible accuracy at {m4['pixel_baseline']['coverage']:.1%} coverage | RGB pixel baseline on synthetic rendered scenes |
| [AI Quality Evaluation](05-AI-Quality-Evaluation/) | 100 cases | {m5['macro_f1']:.3f} macro-F1 | Lexical triage versus provisional project-owned decisions; no safety certification |
| [Search Relevance](06-Search-Relevance/) | 30 queries / 120 results | {m6['ndcg_at_4']:.3f} NDCG@4 | Lexical ranking on fictional judged results, not live search |

## Evidence and provenance

| Evidence | Status |
|---|---|
| Source JSON, normalized JSONL/CSV, SHA-256 manifest | Achieved |
| Runnable shared analysis and per-project entry points | Achieved |
| Executed notebooks with real metric output | Achieved |
| Deterministic PNG charts and 120 rendered images | Achieved |
| Automated schema, span, pixel, range and reproduction tests | Achieved |
| Independent qualified human adjudication | **Pending** |
| Malayalam native review | **Pending** |
| Employer-policy validation or production certification | **Not claimed** |

See [`PROVENANCE.md`](PROVENANCE.md), [`VALIDATION.md`](VALIDATION.md), and [`DATA_MANIFEST.json`](DATA_MANIFEST.json).

## Decision workflow
**Evidence → rubric → independent baseline/QC → decision → review queue → human adjudication → documented revision.**

Automatic checks prioritize work; they do not silently replace reviewers.

## Run locally
Requires Python and dependencies in `requirements.txt`. No network is required.

- `python build_portfolio.py`
- `python analysis.py --project all`
- `python -m unittest discover -s tests`

## Known limits
All six projects are portfolio-scale and predominantly synthetic. Provisional labels may carry AI-generation bias. There is no completed human inter-annotator study, no live-system evaluation, no natural-image benchmark, and no employment-performance claim.

## Career materials
Printable resume and drafting materials are in [`career/`](career/). Research coverage is deliberately limited and source-linked in [`research/hiring-evidence.md`](research/hiring-evidence.md).

## Contact
- Email: [{EMAIL}](mailto:{EMAIL})
- LinkedIn: [{LINKEDIN}]({LINKEDIN})
- GitHub: [{GITHUB}]({GITHUB})

## License
Original code is MIT licensed. AI-assisted synthetic data is offered under CC BY 4.0; see [`DATA_LICENSE.md`](DATA_LICENSE.md).
"""
    write_text(ROOT / "README.md", readme)
    write_text(ROOT / "LICENSE", """MIT License

Copyright (c) 2026 Sinshad PS

Permission is hereby granted, free of charge, to any person obtaining a copy
of this software and associated documentation files (the "Software"), to deal
in the Software without restriction, including without limitation the rights
to use, copy, modify, merge, publish, distribute, sublicense, and/or sell
copies of the Software, and to permit persons to whom the Software is
furnished to do so, subject to the following conditions:

The above copyright notice and this permission notice shall be included in all
copies or substantial portions of the Software.

THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT.
""")
    write_text(ROOT / "DATA_LICENSE.md", f"""# Synthetic Data License

Unless a source extract states otherwise, the synthetic datasets generated for this portfolio are licensed under the [Creative Commons Attribution 4.0 International License](https://creativecommons.org/licenses/by/4.0/).

Attribution: **AI Generalist Evaluation Portfolio, {AUTHOR}, 2026**.

The data is AI-assisted and synthetic. The license does not represent that provisional labels are human validated, factually complete, safe for production, or free from bias. Third-party research extracts remain under their original publishers’ terms and are not relicensed here.
""")
    write_text(ROOT / ".gitignore", """__pycache__/
*.py[cod]
.venv/
venv/
.DS_Store
.ipynb_checkpoints/
~$*.docx
""")
    write_text(ROOT / "CHANGELOG.md", f"""# Changelog

## {TODAY}
- Built six reproducible AI evaluation projects.
- Added normalized datasets, metrics, charts, notebooks, QA queues, and worked cases.
- Added offline portfolio website, career materials, tests, provenance, licenses, and SHA-256 data manifest.
- Documented pending human review and external-validity limits.
""")
    write_text(ROOT / "PROVENANCE.md", f"""# Provenance

- Author attribution: {AUTHOR}.
- Build date: {TODAY}.
- Deterministic seed: {SEED}.
- Portfolio artifacts are AI-assisted synthetic personal projects, not employment deliverables.
- The author has not represented these artifacts as independently personally reviewed.
- LLM author references and judge outputs are known only as Strawberry fast model, same family unspecified. No model/version or temperature was supplied or invented.
- Annotation, Malayalam, safety, and search references are provisional synthetic labels.
- Multimodal gold values are constructed from deterministic rendering geometry.
- No independent human gold, human kappa, qualification promise, job promise, or production certification is claimed.
""")
    write_text(ROOT / "HUMAN_REVIEW.md", f"""# Human Review Plan

## Current status
As of {TODAY}, qualified independent human review remains pending.

## Launch checklist
- [ ] Assign reviewers and record role-relevant qualifications.
- [ ] Blind reviewers to provisional labels where practical.
- [ ] Complete Malayalam native review fields; do not copy AI scores.
- [ ] Review annotation holdout without editing frozen dictionaries.
- [ ] Adjudicate LLM author/judge disagreements.
- [ ] Inspect safety escalation false negatives before false positives.
- [ ] Review search failures with locale/domain context.
- [ ] Sample image geometry and pixel baseline failures.
- [ ] Record review date, reviewer ID, evidence, confidence, and rationale.
- [ ] Resolve disagreements through a named adjudicator.
- [ ] Re-run analysis and tests after accepted corrections.
- [ ] Update provenance and release notes without retroactively hiding failures.

## Unresolved at build date
Reviewer identities, qualification criteria, adjudication outcomes, Malayalam language scores, external benchmark validity, employer-policy mapping, and production launch approval are unresolved.
""")
    write_text(ROOT / "recruiter-guide.md", """# Recruiter Guide

## Suggested five-minute review
1. Start with the website’s six evidence cards.
2. Open one metrics JSON and its chart.
3. Inspect the corresponding review queue and three worked cases.
4. Review `analysis.py` for label-independent baseline logic.
5. Check tests and `VALIDATION.md`.
6. Read the limitations before interpreting any metric.

## What is demonstrated
Structured evaluation design, reproducible Python analysis, rubric writing, dataset QA, error queues, basic visual generation/pixel analysis, local documentation, and honest scope control.

## What is not demonstrated
Employment experience, production deployment, advanced machine learning, independently validated personal accuracy, natural-image VLM performance, or policy certification.
""")
    write_text(ROOT / "project-defensibility-guide.md", """# Project Defensibility Guide

## Core formulas
- Classification accuracy: correct predictions divided by total predictions.
- Per-class F1: harmonic mean of precision and recall; macro-F1 averages class F1 equally.
- Entity exact match: predicted and reference start, end, and label must all match.
- LLM weighted score: 0.4 accuracy + 0.3 instruction following + 0.2 helpfulness + 0.1 clarity.
- NDCG@4: DCG with gain `2^relevance−1` and `log2(rank+1)` discount, divided by ideal DCG.
- MRR: reciprocal rank of the first result with relevance at least 2.
- Pixel baseline coverage: eligible predictions divided by all visual questions.

## Defensible statements
Say “the pipeline measured X against provisional project references,” not “I achieved X% human accuracy.” Say “rendering-derived ground truth” only for visual geometry. Describe corruption labels as a controlled simulation, never human agreement.

## Known failures
Keyword rules miss paraphrases. Regex entities have narrow recall. Unicode checks cannot judge translation meaning. RGB components can merge or count UI pixels. Lexical safety rules over-trigger and miss context. Search overlap misses semantic intent.

## Replication
Run root analysis, inspect metrics and queues, execute tests, then compare data-file hashes with `DATA_MANIFEST.json`. Review code before trusting any metric.

## Honest ownership discussion
The repository is AI-assisted. A defensible interview explanation identifies which artifacts were generated, which computations can be explained and rerun, which labels remain provisional, and what the author still needs to learn.
""")
    write_text(ROOT / "CITATION.cff", f"""cff-version: 1.2.0
message: "If you use this portfolio code or synthetic data, please cite it."
title: "AI Generalist Evaluation Portfolio"
type: software
authors:
  - family-names: "PS"
    given-names: "Sinshad"
date-released: "{TODAY}"
license: MIT
repository-code: "{GITHUB}/AI-Generalist-Portfolio"
url: "{GITHUB}/AI-Generalist-Portfolio"
version: "1.0.0"
""")


def research_docs() -> None:
    research = ROOT / "research"
    research.mkdir(exist_ok=True)

    source_dirs = [ROOT / "research-input", ROOT / "source-data" / "research", ROOT.parent / "research"]
    copied = []
    for src in source_dirs:
        if src.is_dir() and src.resolve() != research.resolve():
            for f in src.rglob("*"):
                if f.is_file():
                    target = research / "extracts" / f.name
                    target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(f, target)
                    copied.append(target.name)

    urls = []
    for f in (research / "extracts").glob("*") if (research / "extracts").exists() else []:
        try:
            text = f.read_text(encoding="utf-8", errors="ignore")
            urls.extend(re.findall(r"https?://[^\s\"'<>\]]+", text))
        except OSError:
            pass

    required = {
        "Alignerr b7b95845": "https://www.alignerr.com/",
        "micro1 81905881": "https://www.micro1.ai/",
        "RWS FAQ6806": "https://www.rws.com/",
        "TELUS 18136963": "https://www.telusdigital.com/careers",
        "Mercor list_AAABnh7mQVTevMoBjeVBfpFJ": "https://mercor.com/",
    }
    linked = {}
    for label, fallback in required.items():
        identifier = label.split()[-1]
        linked[label] = next((u.rstrip(".,)") for u in urls if identifier in u), fallback)

    lines = ["# Hiring Evidence", "",
             f"Research snapshot prepared {TODAY}. Coverage is limited to the supplied extracts and official pages linked below; it is not a complete market survey.",
             "", "## Source-linked findings", ""]
    for label, url in linked.items():
        status = ""
        if label.startswith("micro1"):
            status = " — supplied posting treated as closed; no current opening or eligibility implied."
        elif label.startswith("Mercor"):
            status = " — supplied posting treated as closed and degree-requiring; current student status does not imply eligibility."
        lines.append(f"- **{label}**: [{url}]({url}){status}")
    lines += [
        "",
        "## Interpretation",
        "- Verify role status and eligibility directly on the official site before applying.",
        "- RWS and TELUS links are retained as limited evidence only; no hiring outcome is promised.",
        "- Alignerr evidence is role-specific only if the supplied official extract supports it.",
        "- Closed micro1 and Mercor postings are historical evidence, not application recommendations.",
        "- No claim is made that this portfolio meets any platform’s assessment, location, degree, or experience requirements.",
        "",
        "## Public privacy boundary",
        "Private account histories and platform-specific duplicate-account matters are intentionally excluded from public portfolio and application artifacts.",
        "",
        f"Copied local research extracts: {', '.join(sorted(copied)) if copied else 'none supplied or discoverable at build time'}.",
    ]
    write_text(research / "hiring-evidence.md", "\n".join(lines))


def make_resume(metrics: dict[str, dict[str, Any]]) -> None:
    career = ROOT / "career"
    career.mkdir(exist_ok=True)
    m4 = metrics["04-Multimodal-AI-Evaluation"]
    m6 = metrics["06-Search-Relevance"]

    resume_html = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>{AUTHOR} — Resume</title>
<style>
@page{{size:A4;margin:12mm}}*{{box-sizing:border-box}}body{{font:11px/1.35 Arial,sans-serif;color:#172334;margin:0}}
h1{{font-size:25px;margin:0;color:#09263b}}h2{{font-size:13px;text-transform:uppercase;letter-spacing:.08em;color:#087f78;border-bottom:1px solid #9cc9c5;margin:10px 0 4px}}
p{{margin:3px 0}}ul{{margin:3px 0 5px 17px;padding:0}}li{{margin:2px 0}}.contact{{font-size:10px}}.grid{{display:grid;grid-template-columns:1fr 1fr;gap:0 18px}}.muted{{color:#526171}}a{{color:#086b75;text-decoration:none}}
</style></head><body>
<h1>{AUTHOR}</h1>
<p class="contact"><a href="mailto:{EMAIL}">{EMAIL}</a> · <a href="{LINKEDIN}">LinkedIn</a> · <a href="{GITHUB}">GitHub</a></p>
<h2>Summary</h2>
<p>Computer Science Engineering student building AI-assisted, reproducible evaluation portfolio projects. Comfortable with basic Python, structured annotation QA, rubrics, error analysis, documentation, and local evidence workflows. Seeking opportunities to learn under human review; no production or employment experience is implied.</p>
<div class="grid"><section>
<h2>Skills</h2>
<p><b>Evaluation:</b> rubrics, pairwise review, classification metrics, confusion matrices, macro-F1, NDCG/MRR, review queues</p>
<p><b>Python:</b> BASIC scripting, JSON/CSV processing, Pillow image generation, unittest</p>
<p><b>Data QA:</b> schema checks, exact spans, Unicode QC, provenance, deterministic builds</p>
<h2>Education</h2>
<p><b>JCT Institutions</b><br>Computer Science Engineering · Expected 2027</p>
<h2>Languages</h2>
<p>Malayalam — Native (self-reported)<br>English — Professional (self-reported)</p>
</section><section>
<h2>Tools</h2>
<p>Python, JSON, CSV, Git/GitHub workflows, HTML/CSS/JavaScript, Pillow, python-docx, Jupyter JSON format</p>
<h2>Personal Portfolio Projects</h2>
<ul>
<li>Built six AI-assisted synthetic evaluation workflows covering 50 LLM pairs, 100 annotation messages, 50 Malayalam pairs, 120 rendered images, 100 safety cases, and 30 search queries.</li>
<li>Implemented label-independent keyword, lexical-ranking, safety-triage, and RGB-pixel baselines with machine-readable metrics and failure queues.</li>
<li>Rendered 120 deterministic PNG scenes and evaluated a pixel baseline at {m4['pixel_baseline']['coverage']:.1%} coverage; artificial-image scope documented.</li>
<li>Measured search ranking at NDCG@4 {m6['ndcg_at_4']:.3f} on 120 fictional judgments; not live-search performance.</li>
<li>Added reproducibility tests, SHA-256 data manifest, offline website, executed notebooks, and human-review plans.</li>
</ul>
</section></div>
<h2>Disclosure</h2>
<p class="muted">Projects are AI-assisted synthetic personal portfolio work, not employment. Metrics describe supported workflow evidence, not independently validated personal accuracy. Human validation remains pending.</p>
</body></html>"""
    write_text(career / "resume.html", resume_html)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.38)
    section.bottom_margin = Inches(0.38)
    section.left_margin = Inches(0.48)
    section.right_margin = Inches(0.48)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(8.5)
    styles["Normal"].paragraph_format.space_after = Pt(1.5)

    h = doc.add_heading(AUTHOR, 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.runs[0].font.size = Pt(20)
    c = doc.add_paragraph(f"{EMAIL}  |  {LINKEDIN}  |  {GITHUB}")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].font.size = Pt(8)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = None

    heading("Summary")
    doc.add_paragraph("Computer Science Engineering student building AI-assisted, reproducible evaluation portfolio projects. Basic Python, rubrics, annotation QA, metrics, error analysis, and evidence documentation. Seeking supervised opportunities to learn; no production or employment experience is implied.")
    heading("Skills & Tools")
    doc.add_paragraph("Evaluation: pairwise review, confusion matrices, macro-F1, NDCG/MRR, review queues | Python: BASIC scripting, JSON/CSV, Pillow, unittest | QA: schemas, entity spans, Unicode checks, provenance | Tools: Git/GitHub workflows, HTML/CSS/JavaScript, python-docx")
    heading("Personal Portfolio Projects — not employment")
    bullets = [
        "Six AI-assisted synthetic workflows: 50 LLM pairs, 100 annotation messages, 50 Malayalam pairs, 120 rendered images, 100 safety cases, and 30 search queries.",
        "Implemented keyword, lexical-ranking, safety-triage, and RGB-pixel baselines without consulting evaluation labels during prediction.",
        f"Rendered 120 deterministic PNG scenes; pixel baseline coverage {m4['pixel_baseline']['coverage']:.1%}. Artificial-image limitation documented.",
        f"Measured NDCG@4 {m6['ndcg_at_4']:.3f} on 120 fictional search judgments; not live-search performance.",
        "Added machine-readable metrics, failure queues, executed notebooks, tests, SHA-256 manifest, offline website, and human-review plans.",
    ]
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")
    heading("Education")
    doc.add_paragraph("JCT Institutions — Computer Science Engineering, expected 2027")
    heading("Languages")
    doc.add_paragraph("Malayalam — Native (self-reported) | English — Professional (self-reported)")
    heading("Disclosure")
    doc.add_paragraph("AI-assisted synthetic personal projects, not employment. Metrics are workflow evidence, not independently validated personal accuracy. Human validation is pending.")
    doc.save(career / "resume.docx")

    write_text(career / "LinkedIn.md", f"""# LinkedIn Content — DRAFTS ONLY, NOT PUBLISHED

## Headline draft
Computer Science Engineering Student (Expected 2027) | AI Evaluation & Data QA Portfolio | Basic Python | Malayalam & English

## About draft
I am a Computer Science Engineering student at JCT Institutions, expected to graduate in 2027. I am building an AI-assisted synthetic portfolio to learn evaluation workflows: rubric design, data annotation QA, reproducible metrics, review queues, and honest error analysis.

My current work is personal portfolio practice, not employment or production deployment. I use basic Python and document where human validation is still needed. I am especially interested in supervised AI evaluation and data-quality opportunities where I can continue learning.

## Project descriptions
- **AI evaluation portfolio:** Six local projects with runnable analysis, QA queues, charts, notebooks, tests, and provenance.
- **Synthetic visual evaluation:** 120 Pillow-rendered scenes with geometry-derived answers and a pixel baseline.
- **Search relevance:** 30 fictional queries and 120 judgments evaluated with NDCG@4 and MRR.

## Skills draft
AI evaluation, data annotation, quality assurance, rubric writing, basic Python, JSON, CSV, error analysis, confusion matrices, macro-F1, NDCG, MRR, Malayalam, English.

## Featured draft
Feature the monorepo `{GITHUB}/AI-Generalist-Portfolio`, the offline website screenshot, and the project defensibility guide.

## Recommendations draft
Do not fabricate recommendations. Request a recommendation only from a real educator, collaborator, or reviewer who directly inspected the work.
""")
    write_text(career / "github-profile-README.md", f"""# GitHub Profile README Recommendation

## Profile copy
Hi, I’m {AUTHOR}, a Computer Science Engineering student at JCT Institutions (expected 2027). I am learning AI evaluation, data QA, and basic Python through reproducible personal portfolio projects.

- Malayalam: native (self-reported)
- English: professional (self-reported)
- Current focus: evaluation rubrics, annotation QA, metrics, review workflows
- Contact: {EMAIL}

## Current monorepo
**AI-Generalist-Portfolio** — six AI-assisted synthetic evaluation projects with local analysis, tests, evidence queues, and an offline website.

Recommended description: `Reproducible AI evaluation and data-QA portfolio: six synthetic projects, transparent baselines, review queues, tests, and offline evidence.`

Recommended topics: `ai-evaluation`, `data-annotation`, `quality-assurance`, `python`, `nlp`, `search-relevance`, `multimodal`, `portfolio`

## Pinning
Pin the one existing monorepo. Do not imply that separate project repositories exist unless they are actually created.
""")
    write_text(career / "website-copy.md", """# Website Copy Notes

## Primary message
Reproducible AI evaluation practice with transparent baselines, real computed metrics, visible failure queues, and explicit human-review boundaries.

## Evidence promise
Every project links inputs, methodology, metrics, chart, QA queue, and worked examples.

## Disclosure
AI-assisted synthetic personal projects; not employment, production validation, or independently verified personal accuracy.
""")
    write_text(career / "application-kit.md", f"""# Application Kit

## Short summary
Computer Science Engineering student, expected 2027, with an AI-assisted synthetic portfolio demonstrating basic Python, structured evaluation, data QA, and reproducible evidence workflows.

## Short introduction
Hello, I’m {AUTHOR}. I’m studying Computer Science Engineering at JCT Institutions and expect to graduate in 2027. I have built a local portfolio of six synthetic AI evaluation projects to practice rubric application, annotation quality checks, metrics, and reviewer escalation. I am seeking supervised opportunities where I can learn and contribute carefully.

## Honest interview talking points
- I can explain and rerun each metric in the repository.
- I used AI assistance and do not present the artifacts as independently human validated.
- My Python level is basic; I focused on transparent standard-library implementations.
- I separated provisional labels, automatic baselines, and pending human review.
- I can identify baseline failures and propose calibration-only improvements.
- I still need experience with production systems, team review, advanced ML, and real operational policy.

## Human-review launch checklist
- [ ] Date reviewed: __________
- [ ] Reviewer identity and qualification recorded
- [ ] LLM disagreements adjudicated
- [ ] Malayalam native review completed
- [ ] Annotation holdout remains untouched during calibration
- [ ] Safety false negatives reviewed first
- [ ] Search locale assumptions checked
- [ ] Visual pixel failures sampled
- [ ] Resume metrics checked against current JSON
- [ ] No unresolved claim copied into resume
- [ ] Provenance and changelog updated

Unresolved on {TODAY}: qualified reviewer assignment, completed human scores, adjudication, external validity, role eligibility, and production readiness.
""")


def website(metrics: dict[str, dict[str, Any]]) -> None:
    website_dir = ROOT / "website"
    website_dir.mkdir(exist_ok=True)

    cards = []
    samples = {}
    for dirname, meta in PROJECTS.items():
        m = metrics[dirname]
        short = meta["short"]
        measure = {
            "llm": lambda: f"{m['agreement_rate']:.1%} AI-reference agreement",
            "annotation": lambda: f"{m['holdout']['topic']['accuracy']:.1%} holdout topic accuracy",
            "multilingual": lambda: f"{m['automatic_qc_flags']} automatic QC flags",
            "multimodal": lambda: f"{m['pixel_baseline']['accuracy']:.1%} eligible pixel accuracy / {m['pixel_baseline']['coverage']:.1%} coverage",
            "safety": lambda: f"{m['macro_f1']:.3f} heuristic macro-F1",
            "search": lambda: f"{m['ndcg_at_4']:.3f} NDCG@4",
        }[short]()
        chart_uri = "data:image/png;base64," + base64.b64encode(
            (ROOT / dirname / "results" / "chart.png").read_bytes()
        ).decode("ascii")
        cards.append({
            "id": short, "title": meta["title"], "description": meta["description"],
            "measure": measure, "folder": f"../{dirname}/", "chart": chart_uri,
            "validation": "Rendering-derived visual ground truth" if short == "multimodal"
            else "Provisional/synthetic labels; human validation pending",
        })
        samples[short] = sample_rows(dirname)

    payload = json.dumps({"cards": cards, "samples": samples}, ensure_ascii=False).replace("</", "<\\/")
    page = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{AUTHOR} — AI Evaluation Portfolio</title>
<style>
:root{{--navy:#071827;--panel:#10283a;--panel2:#163348;--ink:#edf5f7;--muted:#a9bdc8;--teal:#23c7b5;--amber:#f5b43c;--line:#29485b}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--navy);color:var(--ink);font:16px/1.55 system-ui,-apple-system,Segoe UI,sans-serif}}
a{{color:var(--teal)}}header,main,footer{{max-width:1180px;margin:auto;padding:24px}}header{{padding-top:55px}}
.eyebrow{{color:var(--amber);text-transform:uppercase;letter-spacing:.15em;font-weight:800;font-size:.78rem}}
h1{{font-size:clamp(2.5rem,7vw,5.4rem);line-height:.98;max-width:950px;margin:.2em 0}}h1 span{{color:var(--teal)}}
.lead{{font-size:1.17rem;color:var(--muted);max-width:850px}}.notice{{border-left:4px solid var(--amber);padding:13px 17px;background:#172a38;margin:24px 0}}
.controls{{display:flex;gap:12px;flex-wrap:wrap;margin:28px 0}}button,select,textarea,input{{font:inherit}}
button,.button{{border:1px solid var(--teal);background:transparent;color:var(--ink);padding:10px 15px;border-radius:8px;cursor:pointer;text-decoration:none}}
button:hover,.button:hover{{background:var(--teal);color:#04151d}}.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(310px,1fr));gap:18px}}
.card{{background:linear-gradient(145deg,var(--panel),var(--panel2));border:1px solid var(--line);border-radius:16px;padding:20px;box-shadow:0 14px 35px #0004}}
.card h2{{font-size:1.35rem;margin:.2em 0}}.metric{{font-size:1.2rem;color:var(--amber);font-weight:800}}.badge{{display:inline-block;border:1px solid var(--line);border-radius:99px;padding:4px 9px;color:var(--muted);font-size:.76rem}}
.card img{{width:100%;border-radius:10px;margin-top:12px}}section{{margin:45px 0}}.review{{display:grid;grid-template-columns:270px 1fr;gap:18px}}
.panel{{background:var(--panel);border:1px solid var(--line);padding:18px;border-radius:14px}}pre{{white-space:pre-wrap;word-break:break-word;color:#d6e5e9}}
label{{display:block;color:var(--muted);margin-top:10px}}select,input,textarea{{width:100%;background:#081c2b;color:var(--ink);border:1px solid var(--line);border-radius:7px;padding:9px}}
textarea{{min-height:90px}}.steps{{display:flex;gap:6px;flex-wrap:wrap}}.steps span{{background:#102b3c;border-radius:99px;padding:7px 11px}}
footer{{border-top:1px solid var(--line);color:var(--muted)}}@media(max-width:720px){{.review{{grid-template-columns:1fr}}header{{padding-top:30px}}}}
@media print{{body{{background:white;color:#111}}.card,.panel{{background:white;box-shadow:none}}button{{display:none}}}}
</style></head><body>
<header><div class="eyebrow">AI evaluation · data quality · evidence review</div>
<div class="eyebrow">Sinshad PS · CSE student · Malayalam + English</div>
<h1>AI evaluation.<br><span>Evidence first.</span></h1>
<p class="lead">A local, reproducible portfolio by {AUTHOR}: six AI-assisted synthetic personal projects with computed metrics, failure queues, worked evidence, and no hidden network dependency.</p>
<div class="notice"><b>Truthful scope:</b> This is not employment work. Provisional AI/synthetic labels are not human gold. Human review remains pending except that visual gold values are deterministically constructed from rendering geometry.</div>
<div class="controls"><a class="button" href="../README.md">Repository guide</a><a class="button" href="../VALIDATION.md">Validation evidence</a><a class="button" href="../career/resume.html">Printable resume</a></div>
</header><main>
<section><div class="eyebrow">About / skills</div><h2>Practical AI review, with traceable decisions.</h2><p class="lead">Computer Science Engineering student, expected graduation 2027. Focused on remote AI generalist, evaluator, annotation and content-quality work. Malayalam and English language skills are self-reported; multilingual benchmark judgments await native review.</p><div class="steps"><span>Python — basic</span><span>Pairwise response evaluation</span><span>Annotation &amp; QA</span><span>Safety triage</span><span>Search intent</span><span>Multimodal review</span></div></section>
<section><div class="eyebrow">Six project cards</div><h2>Measured workflow evidence</h2><div class="grid" id="cards"></div></section>
<section><div class="eyebrow">Decision workflow</div><h2>Evidence before claims</h2><div class="steps"><span>1 Evidence</span><span>2 Rubric</span><span>3 QC</span><span>4 Decision</span><span>5 Queue</span><span>6 Human adjudication</span></div></section>
<section><div class="eyebrow">Interactive local demo</div><h2>Prompt → output → rubric → QC → decision</h2>
<div class="review"><div class="panel"><label>Project<select id="project"></select></label><label>Worked case<select id="case"></select></label>
<p class="badge">Manual review pending</p><button id="export">Export review JSON</button></div>
<div class="panel"><h3 id="sid"></h3><label>Input / source context</label><pre id="input"></pre><img id="scene" alt="Actual rendered evaluation scene" style="display:none;max-width:100%;border-radius:12px"><label>Output</label><pre id="output"></pre><label>Rubric</label><pre id="rubric"></pre><label>QC</label><pre id="qc"></pre><label>Provisional decision</label><pre id="decision"></pre><label>Evidence</label><pre id="evidence"></pre>
<h3>Reviewer scratchpad — local only</h3><label>Score 1–5<input id="score" type="number" min="1" max="5"></label><label>Reviewer decision<select id="rdecision"><option value="">Select</option><option>uphold</option><option>revise</option><option>escalate</option><option>insufficient evidence</option></select></label><label>Feedback<textarea id="feedback"></textarea></label></div></div></section>
<section><div class="eyebrow">Evidence boundary</div><div class="grid"><div class="card"><h2>Achieved pipeline</h2><p>Local data, runnable code, deterministic charts/images, actual computed summaries, review queues, manifest, tests, and documentation.</p></div><div class="card"><h2>Human validation pending</h2><p>No fabricated reviewer ratings, human agreement, certification, production accuracy, or employment performance. Malayalam native review and adjudication remain open.</p></div></div></section>
</main><footer><p>{AUTHOR} · <a href="mailto:{EMAIL}">{EMAIL}</a> · <a href="{LINKEDIN}">LinkedIn</a> · <a href="{GITHUB}">GitHub</a></p><p>AI-assisted synthetic personal projects, not employment.</p></footer>
<script>
const DATA={payload};
const esc=s=>String(s??"");
const cards=document.getElementById("cards");
DATA.cards.forEach(c=>{{const el=document.createElement("article");el.className="card";el.dataset.id=c.id;
el.innerHTML=`<span class="badge">${{esc(c.validation)}}</span><h2>${{esc(c.title)}}</h2><p>${{esc(c.description)}}</p><p class="metric">${{esc(c.measure)}}</p><a href="${{c.folder}}">Open evidence folder</a><img alt="Computed chart for ${{esc(c.title)}}" src="${{c.chart}}">`;cards.appendChild(el)}});
const project=document.getElementById("project"),caseSel=document.getElementById("case");
DATA.cards.forEach(c=>{{let o=document.createElement("option");o.value=c.id;o.textContent=c.title;project.appendChild(o)}});
function fillCases(){{caseSel.innerHTML="";DATA.samples[project.value].forEach((s,i)=>{{let o=document.createElement("option");o.value=i;o.textContent=s.id;caseSel.appendChild(o)}});render()}}
function render(){{const s=DATA.samples[project.value][Number(caseSel.value||0)];["input","output","rubric","qc","decision","evidence"].forEach(k=>document.getElementById(k).textContent=esc(s[k]));document.getElementById("sid").textContent="Case "+s.id;const image=document.getElementById("scene");image.style.display=s.image?"block":"none";if(s.image)image.src=s.image;["score","rdecision","feedback"].forEach(k=>document.getElementById(k).value="")}}
project.onchange=fillCases;caseSel.onchange=render;fillCases();
document.getElementById("export").onclick=()=>{{const s=DATA.samples[project.value][Number(caseSel.value||0)];const scoreField=document.getElementById("score");if(!scoreField.checkValidity()){{scoreField.reportValidity();return}}const out={{project:project.value,case_id:s.id,score:document.getElementById("score").value||null,reviewer_decision:document.getElementById("rdecision").value,feedback:document.getElementById("feedback").value,status:"locally_exported_unverified_review",exported_at:new Date().toISOString()}};const b=new Blob([JSON.stringify(out,null,2)],{{type:"application/json"}});const a=document.createElement("a");a.href=URL.createObjectURL(b);a.download=`review-${{project.value}}-${{s.id}}.json`;a.click();URL.revokeObjectURL(a.href)}};
</script></body></html>"""
    write_text(website_dir / "index.html", page)


TEST_SOURCE = r'''
import csv, hashlib, json, subprocess, sys, unittest
from pathlib import Path
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
def jl(path):
 return [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]
class PortfolioTests(unittest.TestCase):
 def test_schemas_ids_and_counts(self):
  cases={
   "01-LLM-Evaluation/data/pairs.jsonl":50,
   "01-LLM-Evaluation/data/judge-decisions.jsonl":50,
   "02-AI-Data-Annotation/data/production_data.jsonl":100,
   "03-Multilingual-AI-Evaluation/data/pairs.jsonl":50,
   "04-Multimodal-AI-Evaluation/data/visual_answers.jsonl":240,
   "05-AI-Quality-Evaluation/data/cases.jsonl":100,
   "06-Search-Relevance/data/queries.jsonl":30}
  for name,n in cases.items():
   rows=jl(ROOT/name);self.assertEqual(len(rows),n,name)
   key="record_id" if "visual_answers" in name else "id"
   self.assertEqual(len({str(x[key]) for x in rows}),n,name)
 def test_annotation_labels_spans_and_counts(self):
  rows=jl(ROOT/"02-AI-Data-Annotation/data/production_data.jsonl")
  self.assertEqual(CounterLike(rows,"split"),{"calibration":70,"holdout":30})
  for r in rows:
   self.assertIn(r["topic"],{"billing","account","delivery","technical","refund","general"})
   self.assertIn(r["sentiment"],{"positive","neutral","negative","mixed"})
   self.assertIn(r["urgency"],{"low","medium","high"})
   for e in r["entity_spans"]:
    self.assertEqual(r["text"][e["start"]:e["end"]],e["text"])
    self.assertIn(e["label"],{"ORDER_ID","PRODUCT","DATE","ORG"})
 def test_judge_score_range(self):
  for r in jl(ROOT/"01-LLM-Evaluation/data/judge-decisions.jsonl"):
   for side in ("scores_a","scores_b"):
    for key in ("accuracy","instruction_following","helpfulness","clarity"):
     self.assertIsInstance(r[side][key],int);self.assertTrue(1<=r[side][key]<=5)
 def test_image_pixels(self):
  files=sorted((ROOT/"04-Multimodal-AI-Evaluation/data/images").glob("*.png"))
  self.assertEqual(len(files),120)
  hashes=set()
  for p in files:
   im=Image.open(p).convert("RGB");self.assertEqual(im.size,(480,300))
   self.assertGreater(len(im.getcolors(maxcolors=1000000)),2)
   hashes.add(hashlib.sha256(p.read_bytes()).hexdigest())
  self.assertEqual(len(hashes),120)
 def test_output_consistency(self):
  m=json.loads((ROOT/"04-Multimodal-AI-Evaluation/results/metrics.json").read_text())
  self.assertEqual(m["scene_count"],120);self.assertEqual(m["answer_record_count"],240)
  s=json.loads((ROOT/"06-Search-Relevance/results/metrics.json").read_text())
  self.assertEqual(s["query_count"],30);self.assertEqual(s["result_count"],120)
  self.assertTrue(0<=s["ndcg_at_4"]<=1 and 0<=s["mrr_relevance_gte_2"]<=1)
  a=json.loads((ROOT/"02-AI-Data-Annotation/results/metrics.json").read_text())
  for split in ("all","holdout"):
   for field in ("topic","sentiment","urgency"):
    self.assertTrue(0<=a[split][field]["accuracy"]<=1)
    self.assertTrue(0<=a[split][field]["macro_f1"]<=1)
 def test_deterministic_analysis(self):
  paths=[ROOT/d/"results/metrics.json" for d in [
   "01-LLM-Evaluation","02-AI-Data-Annotation","03-Multilingual-AI-Evaluation",
   "04-Multimodal-AI-Evaluation","05-AI-Quality-Evaluation","06-Search-Relevance"]]
  before=[p.read_bytes() for p in paths]
  subprocess.run([sys.executable,str(ROOT/"analysis.py"),"--project","all"],cwd=ROOT,check=True,stdout=subprocess.DEVNULL)
  self.assertEqual(before,[p.read_bytes() for p in paths])
 def test_manifest_hashes(self):
  man=json.loads((ROOT/"DATA_MANIFEST.json").read_text())
  for item in man["files"]:
   p=ROOT/item["path"];self.assertTrue(p.is_file())
   self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(),item["sha256"])
def CounterLike(rows,key):
 out={}
 for r in rows:out[r[key]]=out.get(r[key],0)+1
 return out
if __name__=="__main__":unittest.main()
'''


def tests_and_validation() -> dict[str, Any]:
    write_text(ROOT / "tests" / "test_portfolio.py", TEST_SOURCE)
    result = subprocess.run(
        [sys.executable, "-m", "unittest", "discover", "-s", "tests", "-v"],
        cwd=ROOT, text=True, capture_output=True
    )
    combined = (result.stdout + "\n" + result.stderr).strip()
    validation = {
        "date": TODAY,
        "command": "python -m unittest discover -s tests -v",
        "returncode": result.returncode,
        "passed": result.returncode == 0,
        "output": combined,
    }
    write_text(ROOT / "VALIDATION.md", f"""# Validation Report

Generated automatically on {TODAY}.

- Deterministic seed: `{SEED}`
- Test command: `python -m unittest discover -s tests -v`
- Status: **{"PASS" if result.returncode == 0 else "FAIL"}**

## Covered checks
Schema counts, unique IDs, annotation labels, exact substring spans, judge score ranges, image dimensions and non-background pixels, unique image hashes, metric ranges, output consistency, manifest hashes, and deterministic analysis reproduction.

## Test output
```text
{combined}
```
""")
    if result.returncode != 0:
        raise RuntimeError(f"Portfolio validation failed:\n{combined}")
    return validation


def build_report(manifest: dict[str, Any], validation: dict[str, Any]) -> None:
    report = {
        "author": AUTHOR,
        "build_date": TODAY,
        "seed": SEED,
        "python": sys.version,
        "root": ROOT.name,
        "projects": {
            dirname: read_json(ROOT / dirname / "results" / "metrics.json")
            for dirname in PROJECTS
        },
        "manifest_file_count": len(manifest["files"]),
        "validation": validation,
        "human_review_status": "pending",
        "claims_boundary": "AI-assisted synthetic personal projects, not employment or production certification",
    }
    write_json(ROOT / "BUILD_REPORT.json", report)


def main() -> None:
    os.environ["PYTHONHASHSEED"] = str(SEED)
    source = normalize_inputs()
    prepare_projects(source)
    requirements()
    run_analysis()

    metrics = {
        dirname: read_json(ROOT / dirname / "results" / "metrics.json")
        for dirname in PROJECTS
    }
    for dirname in PROJECTS:
        project_docs(dirname, metrics[dirname])

    create_notebooks()
    root_docs(metrics)
    research_docs()
    make_resume(metrics)
    website(metrics)
    import polish
    polish.run()
    manifest = create_manifest()
    validation = tests_and_validation()
    build_report(manifest, validation)
    print(json.dumps({
        "status": "built",
        "root": str(ROOT),
        "projects": len(PROJECTS),
        "validation": "PASS",
        "data_files": len(manifest["files"]),
    }, indent=2))


if __name__ == "__main__":
    main()
